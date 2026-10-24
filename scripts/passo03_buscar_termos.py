import collections
import csv
import difflib
import hashlib
import json
import os
import re
import sys
import unicodedata

import openpyxl

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import revisao_comum as rc

XLSX = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')
TEXTOS_DIR = os.path.join(AQUI, 'resultado', 'textos')
PPC_DIR = os.path.join(AQUI, 'entradas', 'ppc')
TERMOS_CSV = os.path.join(AQUI, 'entradas', 'termos.csv')
FORA_DO_PPC_CSV = os.path.join(AQUI, 'entradas', 'decisoes', 'paginas_fora_do_ppc.csv')
MANUAIS_CSV = os.path.join(AQUI, 'entradas', 'decisoes', 'mencoes_manuais.csv')
CABECALHO = ['id_ocorrencia', 'id_ppc', 'pagina', 'posicao', 'termo', 'trecho', 'triagem_auto', 'ocr_fraco']
PRIMEIRO_ID = 300000
ALERTAS_CSV = os.path.join(AQUI, 'resultado', 'alertas_busca_aproximada.csv')
CONGELAMENTO_JSON = os.path.join(AQUI, 'resultado', 'congelamento.json')

JANELA = 150
CONFIANCA_MINIMA = 85

RE_BIB = re.compile(r'bibliografia (basica|complementar|recomendada|principal)|bibliografia\s*:|referencias (bibliograficas|basicas|complementares)|^\s*referencias\b|^\s*bibliografia\b', re.M)
RE_RESET = re.compile(r'\bementa\b|\bobjetivos? (geral|gerais|especificos)\b|conteudo programatico|unidade curricular|\bdisciplina\s*[:–-]|componente curricular|codigo:|pre-?requisitos?\s*:|\batividade\s*:|\bdescricao\s*:|\bcomponente\s*:|\bnome\s*:')
RE_INICIO_FICHA = re.compile(r'\bementa\b|\bobjetivos?\b|carga[\s-]*horaria|\bperiodo\b|pre-?requisitos?|\bcreditos?\b|\bcodigo\s*:|\bnatureza\s*:|\bcategoria\s*:|\bch\s*:|\bsemestre\b')
RE_CITACAO = re.compile(r'\bisbn\b|\bissn\b|disponivel em|\bed\. |editora|\bpress\b|\bltc\b|elsevier|bookman|springer|pearson|addison|wiley|mcgraw|prentice|\bp\. ?\d|o.?reilly|manning|packt|apress|mit press|cambridge|oxford|morgan kaufmann|\bcrc|\bsage|new york:|\bet al\.|[a-z .]{3,25}: [a-z&. ]{3,40}, (19|20)\d\d')
RE_AUTOR = re.compile(r'\b[a-z-]{3,}, [a-z]\.\s?(?:[a-z]\.)?')
RE_BIB = re.compile(RE_BIB.pattern.replace(' ', r'\s'), re.M)
RE_CITACAO = re.compile(RE_CITACAO.pattern.replace(' ', r'\s'))
RE_AUTOR = re.compile(RE_AUTOR.pattern.replace(' ', r'\s'))
RE_SUMARIO = re.compile(r'(\. ?){5,}\s*\d+')
RE_ATA = re.compile(r'\bata\b.{0,120}\b(reuni[aã]o|colegiado|nde|conselho)|reuni[aã]o (ordin[aá]ria|extraordin[aá]ria)|aprovad[oa] por unanimidade|\bpauta\b|foi colocad[ao] em (discuss[aã]o|vota[cç][aã]o)', re.I)
RE_EQUIV = re.compile('⇐⇒|⇔|⇒|<=>')
RE_MARCA_OCR = re.compile(r'^\[\[OCR confiança=(\d+)%[^\]]*\]\]\n?')
RE_CONTROLE = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f]')


def dobra(s):
    return ''.join(unicodedata.normalize('NFD', ch)[0] for ch in s)


def espaco(s):
    return re.sub(r'\s+', ' ', RE_CONTROLE.sub('', s)).strip()


def carregar_termos():
    termos = []
    with open(TERMOS_CSV, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f, delimiter=';'):
            termo, sigla = r['termo'].strip(), r['tipo'].strip() == 'sigla'
            t = dobra(termo)
            if not sigla:
                t = t.lower()
            corpo = r'[\s\-]+'.join(re.escape(p) for p in re.split(r'[\s\-]+', t) if p)
            if sigla:
                padrao = r'(?<![A-Za-z0-9])(?<!-)%s(?![A-Za-z0-9])(?!-\d)' % corpo
            else:
                padrao = corpo
            termos.append(dict(termo=termo, sigla=sigla, pat=re.compile(padrao)))
    if not termos:
        raise SystemExit('Nenhum termo em ' + TERMOS_CSV)
    return termos


def carregar_fora_do_ppc():
    intervalos = collections.defaultdict(list)
    if os.path.exists(FORA_DO_PPC_CSV):
        with open(FORA_DO_PPC_CSV, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                if not (r['codigo_curso'] or '').strip():
                    continue
                inicio, fim = int(r['pag_inicio']), int(r['pag_fim'])
                if inicio > fim:
                    raise SystemExit('paginas_fora_do_ppc.csv: pag_inicio maior que pag_fim para o curso ' + r['codigo_curso'])
                intervalos[r['codigo_curso'].strip()].append((inicio, fim))
    return intervalos


def carregar_manuais(termos):
    manuais = collections.defaultdict(list)
    if not os.path.exists(MANUAIS_CSV):
        return manuais
    validos = {t['termo'] for t in termos}
    with open(MANUAIS_CSV, encoding='utf-8-sig', newline='') as f:
        for r in csv.DictReader(f, delimiter=';'):
            if not (r['codigo_curso'] or '').strip():
                continue
            if r['termo'] not in validos:
                raise SystemExit('mencoes_manuais.csv: termo "%s" não está em termos.csv' % r['termo'])
            if not re.search(r'«.+»', r['trecho'] or ''):
                raise SystemExit('mencoes_manuais.csv: o trecho precisa ter a menção entre « » (curso %s, página %s)' % (r['codigo_curso'], r['pagina']))
            manuais[(r['codigo_curso'].strip(), int(r['pagina']))].append((r['termo'], r['trecho'].strip()))
    return manuais


def paginas_do_texto(caminho):
    with open(caminho, encoding='utf-8') as f:
        partes = re.split(r'\[\[PÁGINA (\d+)\]\]\n', f.read())
    saida = []
    for i in range(1, len(partes) - 1, 2):
        corpo, conf = partes[i + 1], None
        m = RE_MARCA_OCR.match(corpo)
        if m:
            conf, corpo = int(m.group(1)), corpo[m.end():]
        saida.append((int(partes[i]), corpo, conf))
    return saida


def secao(fl, s, e, ctx, anterior):
    if RE_SUMARIO.search(ctx):
        return 'sumario'
    antes = anterior + ' ' + fl[:s]
    janela_antes = fl[max(0, s - 150):s]
    resets = list(RE_RESET.finditer(janela_antes))
    if resets:
        janela_antes = janela_antes[resets[-1].end():]
    ficha = RE_INICIO_FICHA.search(fl[e:e + 80])
    if ficha:
        trecho_ate_ficha = fl[e:e + ficha.start()]
        inicio_linha = fl.rfind('\n', 0, s) + 1
        if not (RE_CITACAO.search(trecho_ate_ficha) or RE_AUTOR.search(trecho_ate_ficha) or RE_AUTOR.search(fl[max(0, s - 40):s])
                or re.search(r"[a-z'-]{3,},\s+[a-z]", fl[inicio_linha:s])):
            return 'corpo'
    depois = fl[e:e + 150]
    h = RE_BIB.search(depois)
    if h:
        depois = depois[:h.start()]
    perto = janela_antes + ' ' + depois
    if RE_CITACAO.search(perto) or RE_AUTOR.search(perto):
        return 'bibliografia'
    ultimo_bib = max((m.start() for m in RE_BIB.finditer(antes)), default=-1)
    if ultimo_bib >= 0:
        ultimo_reset = max((m.start() for m in RE_RESET.finditer(antes)), default=-1)
        depois_do_cabecalho = antes[ultimo_bib:][-400:] + ' ' + fl[s:e + 150]
        if ultimo_bib > ultimo_reset and (RE_CITACAO.search(depois_do_cabecalho) or RE_AUTOR.search(depois_do_cabecalho)):
            return 'bibliografia'
    return 'corpo'


def limites_ok(f, s, e):
    antes_ok = s == 0 or not f[s - 1].isalnum() or (f[s - 1].islower() and f[s].isupper())
    depois_ok = e >= len(f) or not f[e].isalnum() or (f[e - 1].islower() and f[e].isupper())
    return antes_ok and depois_ok


def achar_na_pagina(texto, termos, anterior, fora_do_ppc=False):
    f = dobra(texto)
    fl = f.lower()
    achados = []
    for t in termos:
        for m in t['pat'].finditer(f if t['sigla'] else fl):
            if not t['sigla'] and not limites_ok(f, m.start(), m.end()):
                continue
            achados.append((m.start(), m.end(), t))
    achados.sort(key=lambda a: (-(a[1] - a[0]), a[0]))
    mantidos = []
    for s, e, t in achados:
        if not any(s2 < e and s < e2 for s2, e2, _ in mantidos):
            mantidos.append((s, e, t))
    mantidos.sort()
    descartar_junto = set()
    for k in range(len(mantidos) - 1):
        _, e1, t1 = mantidos[k]
        s2, e2, t2 = mantidos[k + 1]
        entre = texto[e1:s2]
        entre_parenteses = re.fullmatch(r'\s*\(\s*', entre) and re.match(r'\s*\)', texto[e2:e2 + 5])
        colados = len(entre) <= 12 and re.fullmatch(r'[ \t]*(?:[-–—:]|c[oó]digo\s*:)?[ \t]*', entre, re.I)
        if t1['sigla'] != t2['sigla'] and (entre_parenteses or colados):
            descartar_junto.add(k if t1['sigla'] else k + 1)
    saida = []
    for k, (s, e, t) in enumerate(mantidos):
        junto = k in descartar_junto
        ctx = espaco(texto[max(0, s - JANELA):e + JANELA])
        trecho = espaco(texto[max(0, s - JANELA):s] + '«' + texto[s:e] + '»' + texto[e:e + JANELA])
        if junto:
            triagem = 'junto com o extenso'
        elif fora_do_ppc:
            triagem = 'fora do PPC'
        else:
            triagem = {'sumario': 'sumário', 'bibliografia': 'bibliografia', 'corpo': ''}[secao(fl, s, e, ctx, anterior)]
            if not triagem:
                antes_c = espaco(texto[max(0, s - 900):s])
                depois_c = espaco(texto[e:e + 500])
                if RE_ATA.search(antes_c[-600:] + ' ' + depois_c[:300]):
                    triagem = 'ata de reunião'
                elif RE_EQUIV.search(antes_c[-120:] + ' ' + depois_c[:120]):
                    triagem = 'tabela de equivalência'
                else:
                    triagem = 'mantida'
        saida.append((s, t['termo'], trecho, triagem))
    return saida


def ppcs_com_texto():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    cab = [c.value for c in next(wb['Dados_Unificados'].iter_rows(max_row=1))]
    i_id, i_cod = cab.index('id_curso'), cab.index('Código do Curso')
    codigo_de = {r[i_id]: str(r[i_cod]).strip() for r in wb['Dados_Unificados'].iter_rows(min_row=2, values_only=True)}
    cab_p = [c.value for c in next(wb['PPC'].iter_rows(max_row=1))]
    j_ppc, j_curso, j_status = cab_p.index('id_ppc'), cab_p.index('id_curso'), cab_p.index('status_coleta')
    por_codigo = collections.defaultdict(list)
    for nome in os.listdir(PPC_DIR):
        if nome.lower().endswith('.pdf'):
            por_codigo[nome.split('_')[0]].append(nome[:-4])
    saida = []
    for r in wb['PPC'].iter_rows(min_row=2, values_only=True):
        if r[j_status] != 1:
            continue
        codigo = codigo_de[r[j_curso]]
        if len(por_codigo[codigo]) != 1:
            raise SystemExit('Curso %s com status_coleta 1 tem %d PDF(s) (esperado 1). Rode os Passos 1 e 2 de novo.' % (codigo, len(por_codigo[codigo])))
        saida.append((r[j_ppc], por_codigo[codigo][0], codigo))
    return sorted(saida)


def sha256_arquivo(caminho):
    with open(caminho, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def gravar_congelamento(ppcs, linhas):
    impressao = {
        'textos': {base: sha256_arquivo(os.path.join(TEXTOS_DIR, base + '.txt')) for _, base, _ in ppcs},
        'paginas_fora_do_ppc': sha256_arquivo(FORA_DO_PPC_CSV) if os.path.exists(FORA_DO_PPC_CSV) else None,
        'termos': sha256_arquivo(TERMOS_CSV),
        'mencoes_manuais': sha256_arquivo(MANUAIS_CSV) if os.path.exists(MANUAIS_CSV) else None,
        'ocorrencia': hashlib.sha256(json.dumps(linhas, ensure_ascii=False).encode('utf-8')).hexdigest(),
        'n_ocorrencias': len(linhas),
    }
    with open(CONGELAMENTO_JSON, 'w', encoding='utf-8') as f:
        json.dump(impressao, f, ensure_ascii=False, indent=1)


def alertas_aproximados(texto, termos_extenso):
    f = dobra(texto).lower()
    exatos = [(m.start(), m.end()) for t in termos_extenso for m in t['pat'].finditer(f)]
    saida = []
    palavras = list(re.finditer(r'[a-z]{4,}', f))
    for t in termos_extenso:
        alvo = ' '.join(re.split(r'[\s\-]+', dobra(t['termo']).lower()))
        for i in range(len(palavras) - 1):
            a, b = palavras[i], palavras[i + 1]
            if b.start() - a.end() > 3 or abs((b.end() - a.start()) - len(alvo)) > 3:
                continue
            candidato = f[a.start():b.end()]
            if candidato == alvo or any(x < b.end() and a.start() < y for x, y in exatos):
                continue
            sim = difflib.SequenceMatcher(None, candidato, alvo).ratio()
            if sim >= 0.85:
                saida.append((t['termo'], texto[a.start():b.end()], round(sim, 2), espaco(texto[max(0, a.start() - 60):b.end() + 60])))
    return saida


def executar():
    for f in (XLSX, TERMOS_CSV):
        if not os.path.exists(f):
            raise SystemExit('Falta %s' % f)
    try:
        open(XLSX, 'r+b').close()
    except PermissionError:
        raise SystemExit('O Excel %s está aberto (ou bloqueado). Feche-o antes de rodar este passo.' % XLSX)
    termos = carregar_termos()
    fora = carregar_fora_do_ppc()
    manuais = carregar_manuais(termos)
    usados_manuais = set()
    ppcs = ppcs_com_texto()
    sem_pdf = rc.codigos_sem_pdf(XLSX)
    de_lado = sum(1 for c in fora if c in sem_pdf) + sum(1 for c, _ in manuais if c in sem_pdf)
    if de_lado:
        print('  decisões de cursos sem PDF em entradas/ppc, deixadas de lado: %d' % de_lado)
    sem_ppc = set(fora) - {codigo for _, _, codigo in ppcs} - sem_pdf
    if sem_ppc:
        raise SystemExit('paginas_fora_do_ppc.csv tem curso que não está entre os PPCs com texto: %s' % sorted(sem_ppc))
    linhas, alertas = [], []
    extensos = [t for t in termos if not t['sigla']]
    for id_ppc, base, codigo in ppcs:
        caminho = os.path.join(TEXTOS_DIR, base + '.txt')
        if not os.path.exists(caminho):
            raise SystemExit('Falta o texto %s: rode o Passo 2.' % caminho)
        anterior = ''
        for pagina, texto, conf in paginas_do_texto(caminho):
            esta_fora = any(a <= pagina <= b for a, b in fora.get(codigo, []))
            for k, (termo_m, trecho_m) in enumerate(manuais.get((codigo, pagina), []), 1):
                usados_manuais.add((codigo, pagina))
                linhas.append([PRIMEIRO_ID + len(linhas), id_ppc, pagina, -k, termo_m, trecho_m, 'fora do PPC' if esta_fora else 'mantida', 0])
            for posicao, termo, trecho, triagem in achar_na_pagina(texto, termos, anterior, esta_fora):
                linhas.append([PRIMEIRO_ID + len(linhas), id_ppc, pagina, posicao, termo, trecho, triagem,
                               int(conf is not None and conf < CONFIANCA_MINIMA)])
            if conf is not None:
                for termo, achado, sim, trecho in alertas_aproximados(texto, extensos):
                    alertas.append([base + '.pdf', pagina, termo, achado, sim, conf, trecho])
            anterior = dobra(texto).lower()[-700:]
    with open(ALERTAS_CSV, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['arquivo', 'pagina', 'termo_parecido', 'texto_encontrado', 'similaridade', 'confianca_ocr', 'trecho'])
        w.writerows(alertas)
    print('  alertas de busca aproximada em páginas de OCR: %d (resultado/alertas_busca_aproximada.csv; só para a validação humana)' % len(alertas))
    sobrou = {m for m in set(manuais) - usados_manuais if m[0] not in sem_pdf}
    if sobrou:
        raise SystemExit('mencoes_manuais.csv: curso/página que não existe entre os textos: %s' % sorted(sobrou))
    chaves = [(l[1], l[2], l[3]) for l in linhas]
    if len(chaves) != len(set(chaves)):
        raise SystemExit('id_ppc + pagina + posicao repetido: erro na busca')
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['OCORRENCIA']
    ws.delete_rows(1, ws.max_row)
    ws.append(CABECALHO)
    for l in linhas:
        ws.append(l)
    ws.freeze_panes = 'A2'
    temporario = XLSX + '.tmp'
    wb.save(temporario)
    os.replace(temporario, XLSX)
    gravar_congelamento(ppcs, linhas)
    resumo(linhas)
    return linhas


def resumo(linhas):
    mantidas = [l for l in linhas if l[6] == 'mantida']
    descartes = collections.Counter(l[6] for l in linhas if l[6] != 'mantida')
    print('  menções: %d (mantidas: %d | descartadas pela máquina: %d)' % (len(linhas), len(mantidas), len(linhas) - len(mantidas)))
    print('  descartes por motivo: %s' % dict(descartes))
    print('  por termo (mantidas): %s' % dict(collections.Counter(l[4] for l in mantidas)))
    print('  mantidas em páginas de OCR fraco: %d | todas as menções em OCR fraco: %d' % (sum(l[7] for l in mantidas), sum(l[7] for l in linhas)))
    print('  PPCs com pelo menos uma menção mantida: %d | PPCs com alguma menção: %d' % (len({l[1] for l in mantidas}), len({l[1] for l in linhas})))
    print('  gravado na aba OCORRENCIA de ' + XLSX)


if __name__ == '__main__':
    executar()
