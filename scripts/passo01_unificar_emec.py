import csv
import glob
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ENTRADA = os.path.join(AQUI, 'entradas', 'emec')
LINKS = os.path.join(AQUI, 'entradas', 'links', 'Tabela_Links_PPCs.csv')
PPC_DIR = os.path.join(AQUI, 'entradas', 'ppc')
SAIDA = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')


CATEGORIAS = [
    (1, 'Macro', 'Corpo do PPC', 'Concepção', 'direta'),
    (2, 'Macro', 'Corpo do PPC', 'Justificativa', 'direta'),
    (3, 'Macro', 'Corpo do PPC', 'Objetivos (curso)', 'direta'),
    (4, 'Macro', 'Corpo do PPC', 'Perfil do egresso', 'direta'),
    (5, 'Macro', 'Corpo do PPC', 'Competências (curso)', 'direta'),
    (6, 'Macro', 'Corpo do PPC', 'Estrutura curricular', 'direta'),
    (7, 'Meso', 'Matriz curricular', 'Componentes curriculares', 'direta'),
    (8, 'Meso', 'Matriz curricular', 'Carga horária', 'direta'),
    (9, 'Meso', 'Matriz curricular', 'Sequência', 'direta'),
    (10, 'Meso', 'Matriz curricular', 'Integração vertical/horizontal', 'direta'),
    (11, 'Meso', 'Matriz curricular', 'Pré-requisitos', 'direta'),
    (12, 'Micro', 'Ficha de disciplina', 'Ementa', 'direta'),
    (13, 'Micro', 'Ficha de disciplina', 'Objetivos (disciplina)', 'direta'),
    (14, 'Micro', 'Ficha de disciplina', 'Competências (disciplina)', 'direta'),
    (15, 'Micro', 'Ficha de disciplina', 'Conteúdos', 'direta'),
    (16, 'Micro', 'Ficha de disciplina', 'Metodologias', 'direta'),
    (17, 'Micro', 'Ficha de disciplina', 'Avaliação', 'direta'),
    (18, 'Relacional', 'Relação entre as partes do PPC', 'Coerência entre macro, meso e micro', 'derivada'),
]


def achar_pdfs(cabecalho, cursos):
    h = [c.strip() for c in cabecalho]
    i_id, i_cod = h.index('id_curso'), h.index('Código do Curso')
    por_codigo = {str(c[i_cod]).strip(): c[i_id] for c in cursos}
    achados = {}
    for nome in sorted(os.listdir(PPC_DIR)):
        if not nome.lower().endswith('.pdf'):
            continue
        codigo = nome.split('_')[0]
        if codigo not in por_codigo:
            raise SystemExit('PDF com código de curso que não está na lista de cursos: ' + nome)
        id_curso = por_codigo[codigo]
        if id_curso in achados:
            raise SystemExit('Dois PDFs para o mesmo curso (%s): %s e %s' % (codigo, achados[id_curso], nome))
        achados[id_curso] = nome
    return achados


def ano_do_pdf(nome):
    ano = nome[:-4].split('_')[-1]
    if not (ano.isdigit() and len(ano) == 4):
        raise SystemExit('PDF sem ano de 4 dígitos no fim do nome: ' + nome)
    return int(ano)


def montar_ppc(cabecalho, cursos):
    h = [c.strip() for c in cabecalho]
    i_id, i_cod = h.index('id_curso'), h.index('Código do Curso')
    links = {}
    with open(LINKS, encoding='utf-8-sig', newline='') as f:
        for l in csv.DictReader(f, delimiter=';'):
            codigo = l['Codigo_Curso'].strip()
            if codigo in links:
                raise SystemExit('Links duplicados para o curso ' + codigo)
            links[codigo] = l
    com_pdf = achar_pdfs(cabecalho, cursos)
    linhas, usadas = [], set()
    for n, c in enumerate(cursos):
        codigo = str(c[i_cod]).strip()
        l = links.get(codigo)
        if l is None:
            raise SystemExit('Curso %s (código %s) sem linha em Tabela_Links_PPCs.csv' % (c[i_id], codigo))
        usadas.add(codigo)
        link = l['Link_download_PPC'].strip()
        tem_link = bool(link) and not link.startswith('(')
        linhas.append([200000 + n, c[i_id], link if tem_link else '', l['Pagina_institucional'].strip(),
                       l['Pagina_do_curso'].strip(), 1 if c[i_id] in com_pdf else 0,
                       ano_do_pdf(com_pdf[c[i_id]]) if c[i_id] in com_pdf else None])
    sobra = set(links) - usadas
    if sobra:
        raise SystemExit('Links sem curso: %s' % sorted(sobra))
    return linhas


COLUNAS_TAMANHO = ['n_paginas_pdf', 'n_paginas_validas', 'n_paginas_ocr', 'n_paginas_pendentes', 'n_caracteres']


def proteger_trabalho_existente():
    if '--recriar' in sys.argv or not os.path.exists(SAIDA):
        return
    try:
        wb = openpyxl.load_workbook(SAIDA, read_only=True)
    except Exception as e:
        raise SystemExit('Não foi possível abrir o Excel atual (%s). Feche-o no Excel; para recriar mesmo assim, use --recriar.' % e)
    com_dados = [aba for aba in ('OCORRENCIA', 'CLASSIFICACAO', 'FATO_ANALISE', 'NIVEL', 'RELACIONAL') if aba in wb.sheetnames and wb[aba].max_row > 1]
    if 'PPC' in wb.sheetnames:
        cab = [c.value for c in next(wb['PPC'].iter_rows(max_row=1))]
        if 'n_paginas_pdf' in cab and any(r[cab.index('n_paginas_pdf')] is not None for r in wb['PPC'].iter_rows(min_row=2, values_only=True)):
            com_dados.append('PPC (colunas de tamanho)')
    wb.close()
    if com_dados:
        raise SystemExit('Rodar o Passo 1 apaga o trabalho dos passos seguintes (%s). Se é isso mesmo que você quer, use --recriar '
                         'e depois rode de novo os Passos 2, 3 e seguintes.' % ', '.join(com_dados))


def executar():
    cabecalho, vistas, unicas, lidas = None, set(), [], 0
    arquivos = sorted(glob.glob(os.path.join(ENTRADA, '*.csv')))
    if not arquivos:
        raise SystemExit('Nenhum CSV do e-MEC em ' + ENTRADA)
    for a in arquivos:
        with open(a, encoding='utf-8-sig', newline='') as f:
            leitor = csv.reader(f, delimiter=';')
            cab = [c.strip() for c in next(leitor)]
            while cab and not cab[-1]:
                cab.pop()
            if cabecalho is None:
                cabecalho = cab
            elif cab != cabecalho:
                raise SystemExit('Cabeçalho diferente em ' + os.path.basename(a))
            n = 0
            for linha in leitor:
                if not any(x.strip() for x in linha):
                    continue
                if len(linha) > len(cab) and any(x.strip() for x in linha[len(cab):]):
                    raise SystemExit('Linha com mais campos que o cabeçalho em %s: %s' % (os.path.basename(a), linha[:3]))
                linha = [x.strip() for x in linha[:len(cab)]]
                if len(linha) != len(cab):
                    raise SystemExit('Linha com menos campos que o cabeçalho em %s: %s' % (os.path.basename(a), linha[:3]))
                n += 1
                if tuple(linha) not in vistas:
                    vistas.add(tuple(linha))
                    unicas.append(linha)
        lidas += n
        print('  %s: %d linhas' % (os.path.basename(a), n))
    i_cat = [c.strip() for c in cabecalho].index('Categoria Administrativa')
    n_unicas = len(unicas)
    unicas = [l for l in unicas if l[i_cat].strip().startswith('Pública')]
    print('  não públicas removidas: %d' % (n_unicas - len(unicas)))
    i_ies = [c.strip() for c in cabecalho].index('Código da IES')
    i_curso = [c.strip() for c in cabecalho].index('Código do Curso')
    unicas.sort(key=lambda l: (int(l[i_ies]), int(l[i_curso])))
    cabecalho = ['id_curso'] + cabecalho
    unicas = [[100000 + n] + l for n, l in enumerate(unicas)]
    proteger_trabalho_existente()
    os.makedirs(os.path.dirname(SAIDA), exist_ok=True)
    ppc = montar_ppc(cabecalho, unicas)
    wb = openpyxl.Workbook()
    abas = [('Dados_Unificados', cabecalho, unicas),
            ('PPC', ['id_ppc', 'id_curso', 'link_ppc', 'pagina_institucional', 'pagina_do_curso', 'status_coleta', 'ano_ppc'] + COLUNAS_TAMANHO,
             [l + [None] * len(COLUNAS_TAMANHO) for l in ppc]),
            ('CATEGORIA', ['id_categoria', 'nivel', 'unidade_documental', 'nome_categoria', 'tipo'], CATEGORIAS),
            ('OCORRENCIA', ['id_ocorrencia', 'id_ppc', 'pagina', 'posicao', 'termo', 'trecho', 'triagem_auto', 'ocr_fraco'], []),
            ('CLASSIFICACAO', ['id_ocorrencia', 'codigo_curso', 'pagina', 'posicao', 'estado', 'nivel', 'id_categoria', 'motivo', 'componente', 'nota'], []),
            ('FATO_ANALISE', ['id_ppc', 'id_categoria', 'n_mencoes', 'n_paginas'], []),
            ('NIVEL', ['id_ppc', 'nivel', 'n_mencoes', 'n_paginas', 'n_sem_topico'], []),
            ('RELACIONAL', ['id_ppc', 'tem_macro', 'tem_meso', 'tem_micro', 'n_niveis'], [])]
    wb.remove(wb.active)
    for nome, cab, linhas in abas:
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for linha in linhas:
            ws.append(list(linha))
        ws.freeze_panes = 'A2'
    temporario = SAIDA + '.tmp'
    wb.save(temporario)
    os.replace(temporario, SAIDA)
    print('  PPC: %d linhas (com PDF: %d, sem PDF: %d) | CATEGORIA: %d | OCORRENCIA, CLASSIFICACAO, FATO_ANALISE, NIVEL e RELACIONAL: só cabeçalho' % (
        len(ppc), sum(r[5] for r in ppc), sum(1 - r[5] for r in ppc), len(CATEGORIAS)))
    print('  linhas lidas: %d | duplicadas removidas: %d | linhas finais: %d' % (lidas, lidas - n_unicas, len(unicas)))
    print('  gravado em ' + SAIDA)
    return unicas


if __name__ == '__main__':
    executar()
