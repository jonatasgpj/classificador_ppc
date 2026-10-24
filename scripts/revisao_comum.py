import collections
import csv
import hashlib
import json
import os

KIT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TABELA = os.path.join(KIT, 'entradas', 'classificacao', 'classificacao.csv')
COLUNAS = ['chave', 'assinatura', 'termo', 'decisao', 'nivel', 'id_categoria', 'motivo', 'componente', 'nota']
MOTIVOS = [('fora_do_ppc', 'Fora do PPC (anexos)'), ('sumario', 'Sumário/índice'), ('bibliografia', 'Bibliografia/referência'), ('ata', 'Ata de reunião'),
           ('equivalencia', 'Tabela de equivalência'), ('docente', 'Lista/quadro de docentes'), ('linha_pesquisa', 'Linha/grupo de pesquisa'),
           ('infraestrutura', 'Laboratório/infraestrutura'), ('sigla_outro_sentido', 'Sigla com outro sentido (não é IA)'),
           ('termo_citado_norma', 'Termo só citado de norma/parecer'), ('duplicata_extenso', 'Sigla colada a um termo por extenso já contado'),
           ('outro_nao_curricular', 'Outro não curricular (exige nota)')]
MOTIVOS_SET = {m for m, _ in MOTIVOS}
TIPOS_EMENTA = [('ementa', 'Ementa propriamente dita'), ('titulo_ficha', 'Título da ficha'), ('cabecalho_ficha', 'Cabeçalho da ficha (código, natureza, pré-requisito)')]
TIPOS_EMENTA_SET = {t for t, _ in TIPOS_EMENTA}
NIVEIS = ('Macro', 'Meso', 'Micro')
DECISOES = {'conta', 'sem_topico', 'descarta', 'duvida'}
FINAIS = ('conta', 'sem_topico', 'descarta')
ID_EMENTA = 12


class ErroTabela(Exception):
    pass


class ErroFila(Exception):
    pass


def _sha256(caminho):
    with open(caminho, 'rb') as f:
        return hashlib.sha256(f.read()).hexdigest()


def codigos_sem_pdf(xlsx):
    """Códigos dos cursos da lista do e-MEC sem PDF em entradas/ppc (status_coleta 0). Quando se roda
    só parte dos PPCs, as decisões e a classificação desses cursos ficam de lado, sem erro."""
    import openpyxl
    wb = openpyxl.load_workbook(xlsx, read_only=True)
    cab = [c.value for c in next(wb['Dados_Unificados'].iter_rows(max_row=1))]
    i_id, i_cod = cab.index('id_curso'), cab.index('Código do Curso')
    codigo_de = {r[i_id]: str(r[i_cod]).strip() for r in wb['Dados_Unificados'].iter_rows(min_row=2, values_only=True)}
    cab_p = [c.value for c in next(wb['PPC'].iter_rows(max_row=1))]
    j_curso, j_status = cab_p.index('id_curso'), cab_p.index('status_coleta')
    saida = {codigo_de[r[j_curso]] for r in wb['PPC'].iter_rows(min_row=2, values_only=True) if r[j_status] != 1}
    wb.close()
    return saida


def verificar_fila_atual(rev_dir, congelamento_json):
    manifesto = json.load(open(os.path.join(rev_dir, 'manifesto.json'), encoding='utf-8'))
    ref = json.load(open(congelamento_json, encoding='utf-8'))
    if manifesto.get('sha256_ocorrencia') != ref['ocorrencia']:
        raise ErroFila('a lista de ocorrências é de outra aba OCORRENCIA que a do Passo 3 atual: rode o passo 4A de novo')
    for nome, h in (manifesto.get('sha256') or {}).items():
        caminho = os.path.join(rev_dir, nome)
        if not os.path.exists(caminho) or _sha256(caminho) != h:
            raise ErroFila('%s mudou ou sumiu (o hash não bate com o manifesto): rode o passo 4A de novo' % nome)
    if not manifesto.get('sha256'):
        raise ErroFila('o manifesto não tem os hashes da lista: rode o passo 4A de novo')
    return manifesto


def _ordem(chave):
    try:
        return tuple(int(x) for x in chave.split('|'))
    except ValueError:
        return (float('inf'), chave)


def ler_tabela(caminho=TABELA):
    if not os.path.exists(caminho):
        return {}
    with open(caminho, encoding='utf-8-sig', newline='') as f:
        leitor = csv.reader(f, delimiter=';')
        cab = next(leitor, None)
        if cab != COLUNAS:
            raise ErroTabela('o cabeçalho de %s não é o esperado (%s)' % (caminho, ';'.join(COLUNAS)))
        saida = {}
        for n, linha in enumerate(leitor, 2):
            if not any(c.strip() for c in linha):
                continue
            if len(linha) != len(COLUNAS):
                raise ErroTabela('linha %d de %s tem %d colunas (esperado %d)' % (n, caminho, len(linha), len(COLUNAS)))
            reg = dict(zip(COLUNAS, linha))
            if reg['chave'] in saida:
                raise ErroTabela('a chave %s aparece duas vezes em %s' % (reg['chave'], caminho))
            saida[reg['chave']] = reg
    return saida


def gravar_tabela(linhas, caminho=TABELA):
    os.makedirs(os.path.dirname(caminho), exist_ok=True)
    temporario = caminho + '.tmp'
    with open(temporario, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(COLUNAS)
        for chave in sorted(linhas, key=_ordem):
            w.writerow([str(linhas[chave].get(c) if linhas[chave].get(c) is not None else '') for c in COLUNAS])
        f.flush()
        os.fsync(f.fileno())
    os.replace(temporario, caminho)


def religar(classificacao, itens):
    na_lista = {i['chave'] for i in itens}
    por_assinatura = collections.defaultdict(list)
    for i in itens:
        por_assinatura[(i['codigo_curso'], i['termo'], i['assinatura'])].append(i['chave'])
    candidatas = {}
    for chave, reg in classificacao.items():
        if chave in na_lista:
            continue
        cand = por_assinatura.get((chave.split('|')[0], reg.get('termo'), reg.get('assinatura')), [])
        candidatas[chave] = cand[0] if (len(cand) == 1 and cand[0] not in classificacao) else None
    disputados = collections.Counter(c for c in candidatas.values() if c)
    saida, religadas, orfas = {}, [], []
    for chave, reg in classificacao.items():
        if chave in na_lista:
            saida[chave] = reg
        elif candidatas[chave] and disputados[candidatas[chave]] == 1:
            saida[candidatas[chave]] = dict(reg, chave=candidatas[chave])
            religadas.append((chave, candidatas[chave]))
        else:
            orfas.append(chave)
    return saida, religadas, orfas


def validar(reg, categorias):
    d = reg.get('decisao')
    if d not in DECISOES:
        return 'decisão inválida: %r' % d
    if d in ('conta', 'sem_topico') and reg.get('nivel') not in NIVEIS:
        return 'nível inválido: %r' % reg.get('nivel')
    if d == 'conta':
        try:
            cat = categorias[int(reg.get('id_categoria'))]
        except (TypeError, ValueError, KeyError):
            return 'id_categoria inválido: %r' % reg.get('id_categoria')
        if cat['tipo'] != 'direta':
            return 'categoria derivada (Relacional) não pode ser atribuída a uma ocorrência'
        if cat['nivel'] != reg['nivel']:
            return 'id_categoria %s é de outro nível (%s) e não de %s' % (reg['id_categoria'], cat['nivel'], reg['nivel'])
        tipo = (reg.get('motivo') or '').strip()
        if tipo and (tipo not in TIPOS_EMENTA_SET or int(reg['id_categoria']) != ID_EMENTA):
            return 'em "conta" o campo motivo só aceita o tipo de ementa (%s), e só para Micro > Ementa' % ', '.join(sorted(TIPOS_EMENTA_SET))
    if d == 'descarta':
        if reg.get('motivo') not in MOTIVOS_SET:
            return 'motivo fora da lista fechada: %r' % reg.get('motivo')
        if reg['motivo'] == 'outro_nao_curricular' and not (reg.get('nota') or '').strip():
            return 'motivo outro_nao_curricular exige nota'
    if d == 'duvida' and not (reg.get('nota') or '').strip():
        return 'dúvida exige nota'
    return None
