import collections
import csv
import json
import os
import sys

import openpyxl

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import passo04a_preparar_fila as fila_mod
import revisao_comum as rc

XLSX = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')
REV = os.path.join(AQUI, 'resultado', 'revisao')
PENDENCIAS_EXTRACAO = os.path.join(AQUI, 'resultado', 'pendencias_extracao.csv')
NIVEIS = rc.NIVEIS
CAB_CLASSIFICACAO = ['id_ocorrencia', 'codigo_curso', 'pagina', 'posicao', 'estado', 'nivel', 'id_categoria', 'motivo', 'componente', 'nota']
CAB_FATO = ['id_ppc', 'id_categoria', 'n_mencoes', 'n_paginas']
CAB_NIVEL = ['id_ppc', 'nivel', 'n_mencoes', 'n_paginas', 'n_sem_topico']
CAB_RELACIONAL = ['id_ppc', 'tem_macro', 'tem_meso', 'tem_micro', 'n_niveis']
ESTADOS = {'conta': 'conta', 'sem_topico': 'sem_topico', 'descarta': 'descartada', 'duvida': 'duvida'}


class ErroConsolidacao(Exception):
    pass


def linhas_classificacao(classificacao, por_chave):
    saida = []
    for chave, reg in classificacao.items():
        i = por_chave[chave]
        d = reg['decisao']
        saida.append([i['id_ocorrencia'], i['codigo_curso'], i['pagina'], i['posicao'], ESTADOS[d],
                      reg['nivel'] if d in ('conta', 'sem_topico') else None,
                      int(reg['id_categoria']) if d == 'conta' else None,
                      reg['motivo'] if (d == 'descarta' or (d == 'conta' and reg.get('motivo'))) else None,
                      reg.get('componente') or None, reg.get('nota') or None])
    saida.sort(key=lambda r: r[0])
    return saida


def tabelas_derivadas(classificacao, por_chave, ppcs_com_pdf, categorias_diretas):
    n_men = collections.Counter()
    pags = collections.defaultdict(set)
    nivel_men, nivel_pags, nivel_sem = collections.Counter(), collections.defaultdict(set), collections.Counter()
    for chave, reg in classificacao.items():
        i = por_chave[chave]
        if reg['decisao'] == 'conta':
            k = (i['id_ppc'], int(reg['id_categoria']))
            n_men[k] += 1
            pags[k].add(i['pagina'])
        if reg['decisao'] in ('conta', 'sem_topico'):
            k = (i['id_ppc'], reg['nivel'])
            nivel_men[k] += 1
            nivel_pags[k].add(i['pagina'])
            if reg['decisao'] == 'sem_topico':
                nivel_sem[k] += 1
    fato = [[p, c, n_men[(p, c)], len(pags[(p, c)])] for p in ppcs_com_pdf for c, _ in categorias_diretas]
    nivel = [[p, n, nivel_men[(p, n)], len(nivel_pags[(p, n)]), nivel_sem[(p, n)]] for p in ppcs_com_pdf for n in NIVEIS]
    relacional = []
    for p in ppcs_com_pdf:
        tem = [int(nivel_men[(p, n)] > 0) for n in NIVEIS]
        relacional.append([p] + tem + [sum(tem)])
    return fato, nivel, relacional


def pendentes(itens, classificacao):
    saida = []
    for i in itens:
        reg = classificacao.get(i['chave'])
        if reg is None:
            saida.append((i, 'sem classificação'))
        elif reg['decisao'] == 'duvida':
            saida.append((i, 'dúvida'))
    return saida


def contar_pendencias_extracao():
    if not os.path.exists(PENDENCIAS_EXTRACAO):
        return 0
    with open(PENDENCIAS_EXTRACAO, encoding='utf-8-sig', newline='') as f:
        linhas = [l for l in csv.reader(f, delimiter=';') if l and any(c.strip() for c in l)]
    return max(0, len(linhas) - 1)


def gravar_aba(wb, nome, cabecalho, linhas):
    ws = wb[nome]
    ws.delete_rows(1, ws.max_row)
    ws.append(cabecalho)
    for l in linhas:
        ws.append(l)
        for cel in ws[ws.max_row]:
            if isinstance(cel.value, str) and cel.value[:1] in ('=', '+', '-', '@'):
                cel.data_type = 's'
    ws.freeze_panes = 'A2'


def _ocorrencias(caminho_xlsx):
    wb = openpyxl.load_workbook(caminho_xlsx, read_only=True)
    cab = [c.value for c in next(wb['OCORRENCIA'].iter_rows(max_row=1))]
    saida = [dict(zip(cab, r)) for r in wb['OCORRENCIA'].iter_rows(min_row=2, values_only=True)]
    wb.close()
    return saida


def executar(parcial=False, aceitar_orfas=False):
    fila_mod.verificar_congelamento(_ocorrencias(XLSX))
    try:
        rc.verificar_fila_atual(REV, fila_mod.CONGELAMENTO_JSON)
    except rc.ErroFila as e:
        raise ErroConsolidacao('lista de ocorrências desatualizada: %s' % e)
    itens = json.load(open(os.path.join(REV, 'fila.json'), encoding='utf-8'))
    por_chave = {i['chave']: i for i in itens}
    wb = openpyxl.load_workbook(XLSX)
    cats = {r[0]: {'nivel': r[1], 'tipo': r[4]} for r in wb['CATEGORIA'].iter_rows(min_row=2, values_only=True)}
    diretas = [(c, v['nivel']) for c, v in sorted(cats.items()) if v['tipo'] == 'direta']
    cab_p = [c.value for c in wb['PPC'][1]]
    ppcs_com_pdf = [r[cab_p.index('id_ppc')] for r in wb['PPC'].iter_rows(min_row=2, values_only=True) if r[cab_p.index('status_coleta')] == 1]

    try:
        classificacao, religadas, orfas = rc.religar(rc.ler_tabela(), itens)
    except rc.ErroTabela as e:
        raise ErroConsolidacao('tabela de classificação inválida: %s' % e)
    sem_pdf = rc.codigos_sem_pdf(XLSX)
    de_lado = [c for c in orfas if c.split('|')[0] in sem_pdf]
    orfas = [c for c in orfas if c.split('|')[0] not in sem_pdf]
    if orfas and not aceitar_orfas:
        raise ErroConsolidacao('%d linha(s) ÓRFÃ(S): a chave não existe mais na lista e a assinatura não acha um único candidato (ex.: %s). '
                               'Reveja ou use --aceitar-orfas (elas ficam de fora).' % (len(orfas), orfas[:3]))
    erros = [(c, rc.validar(reg, cats)) for c, reg in classificacao.items() if rc.validar(reg, cats)]
    if erros:
        raise ErroConsolidacao('%d linha(s) inválida(s) na tabela; primeira: %s -> %s' % (len(erros), erros[0][0], erros[0][1]))
    pend = pendentes(itens, classificacao)
    n_extracao = contar_pendencias_extracao()
    bloqueios = []
    if pend:
        bloqueios.append('%d ocorrência(s) sem classificação ou em dúvida' % len(pend))
    if n_extracao:
        bloqueios.append('%d pendência(s) de extração aberta(s) (resultado/pendencias_extracao.csv)' % n_extracao)

    linhas = linhas_classificacao(classificacao, por_chave)
    gravar_aba(wb, 'CLASSIFICACAO', CAB_CLASSIFICACAO, linhas)
    gravou_derivadas = (not bloqueios) or parcial
    if gravou_derivadas:
        fato, nivel, relacional = tabelas_derivadas(classificacao, por_chave, ppcs_com_pdf, diretas)
        if parcial and bloqueios:
            gravar_aba(wb, 'FATO_ANALISE', CAB_FATO + ['parcial'], [l + [1] for l in fato])
        else:
            gravar_aba(wb, 'FATO_ANALISE', CAB_FATO, fato)
        gravar_aba(wb, 'NIVEL', CAB_NIVEL, nivel)
        gravar_aba(wb, 'RELACIONAL', CAB_RELACIONAL, relacional)
    else:
        gravar_aba(wb, 'FATO_ANALISE', CAB_FATO, [])
        gravar_aba(wb, 'NIVEL', CAB_NIVEL, [])
        gravar_aba(wb, 'RELACIONAL', CAB_RELACIONAL, [])
    temporario = XLSX + '.tmp'
    wb.save(temporario)
    wb.close()
    os.replace(temporario, XLSX)

    with open(os.path.join(REV, 'pendentes.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['chave', 'motivo', 'termo', 'pagina', 'sigla'])
        for i, motivo in pend:
            w.writerow([i['chave'], motivo, i['termo'], i['pagina'], i['sigla']])
    contagem = collections.Counter(c[4] for c in linhas)
    resumo = {'itens': len(itens), 'classificadas': len(classificacao), 'pendentes': len(pend), 'estados': dict(contagem),
              'por_nivel': dict(collections.Counter(c[5] for c in linhas if c[5])), 'pendencias_extracao': n_extracao,
              'bloqueios': bloqueios, 'derivadas_gravadas': gravou_derivadas, 'parcial': bool(parcial and bloqueios),
              'religadas': religadas, 'orfas': orfas, 'de_cursos_sem_pdf': len(de_lado)}
    with open(os.path.join(REV, 'consolidacao.json'), 'w', encoding='utf-8', newline='\n') as f:
        json.dump(resumo, f, ensure_ascii=False, indent=1)
    imprimir(resumo)
    return resumo


def imprimir(r):
    print('  ocorrências: %d | classificadas: %d | sem classificação ou em dúvida: %d' % (r['itens'], r['classificadas'], r['pendentes']))
    print('  estados: %s | por nível: %s' % (r['estados'], r['por_nivel']))
    if r['religadas']:
        print('  linhas religadas pela assinatura: %d' % len(r['religadas']))
    if r['de_cursos_sem_pdf']:
        print('  classificações de cursos sem PDF em entradas/ppc, deixadas de lado: %d' % r['de_cursos_sem_pdf'])
    if r['orfas']:
        print('  ATENÇÃO: %d linha(s) órfã(s) deixadas de fora' % len(r['orfas']))
    if r['derivadas_gravadas'] and not r['parcial']:
        print('  FATO_ANALISE, NIVEL e RELACIONAL gravadas (classificação completa).')
    elif r['derivadas_gravadas']:
        print('  FATO_ANALISE, NIVEL e RELACIONAL gravadas em modo PARCIAL (--parcial): %s' % '; '.join(r['bloqueios']))
    else:
        print('  FATO_ANALISE, NIVEL e RELACIONAL NÃO foram gravadas e ficaram só com o cabeçalho (ainda restam: %s). Lista em resultado/revisao/pendentes.csv.' % '; '.join(r['bloqueios']))
    print('  CLASSIFICACAO gravada com as ocorrências classificadas.')


if __name__ == '__main__':
    try:
        executar(parcial='--parcial' in sys.argv, aceitar_orfas='--aceitar-orfas' in sys.argv)
    except ErroConsolidacao as e:
        raise SystemExit('Consolidação interrompida: %s' % e)
