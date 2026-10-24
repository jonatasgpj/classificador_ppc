import collections
import concurrent.futures
import csv
import glob
import os
import re
import statistics
import subprocess
import sys
import tempfile
import unicodedata

import openpyxl
import pymupdf

pymupdf.TOOLS.mupdf_display_errors(False)

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import revisao_comum as rc
XLSX = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')
PPC_DIR = os.path.join(AQUI, 'entradas', 'ppc')
TEXTOS_DIR = os.path.join(AQUI, 'resultado', 'textos')
PENDENCIAS_DIR = os.path.join(TEXTOS_DIR, 'pendencias')
PENDENCIAS_CSV = os.path.join(AQUI, 'resultado', 'pendencias_extracao.csv')
SUSPEITAS_CSV = os.path.join(AQUI, 'resultado', 'paginas_ocr_suspeitas.csv')
FORA_DO_PPC_CSV = os.path.join(AQUI, 'entradas', 'decisoes', 'paginas_fora_do_ppc.csv')
CONFERIDAS_CSV = os.path.join(AQUI, 'entradas', 'decisoes', 'pendencias_conferidas.csv')
CONTROLE_CSV = os.path.join(AQUI, 'entradas', 'decisoes', 'caracteres_controle.csv')
TESSERACT = os.path.join(r'C:\Program Files\Tesseract-OCR', 'tesseract.exe')
TESSDATA = os.path.join(AQUI, 'entradas', 'ocr')
OCR_PSM = '4'
PDFS_AO_MESMO_TEMPO = 2

RX_CONTROLE = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
LIGADURAS = {'ﬁ': 'fi', 'ﬂ': 'fl', 'ﬀ': 'ff', 'ﬃ': 'ffi', 'ﬄ': 'ffl', 'ﬅ': 'st', 'ﬆ': 'st'}
ACENTOS = {'´': '́', '`': '̀', 'ˆ': '̂', '˜': '̃', '¨': '̈'}
PREFIXOS = {'pré', 'pós', 'ex', 'vice', 'anti', 'semi', 'inter', 'intra', 'ultra'}
CLITICOS = {'se', 'lo', 'la', 'los', 'las', 'lhe', 'lhes', 'me', 'te', 'nos', 'vos'}


LIMITE_TEXTO_OCR = 400
CONFIANCA_SUSPEITA = 85
LIMITE_ILEGIVEL = 0.3
_C = r'[\x00-\x08\x0b\x0c\x0e-\x1f]'
RX_FRAGMENTO = re.compile(r'([A-Za-zÀ-ÿ]*)(' + _C + r'+)([A-Za-zÀ-ÿ]*)')
RX_SIGLA_QUEBRADA = re.compile(r'(?<![A-Za-z])(?:' + _C + r')A(?![A-Za-z])|(?<![A-Za-z])I(?:' + _C + r')(?![A-Za-z])')
PALAVRAS_TERMO = ('artificial', 'intelligence', 'inteligencia')


def termo_quebrado(texto):
    for m in RX_FRAGMENTO.finditer(texto):
        esq = unicodedata.normalize('NFD', m.group(1)).encode('ascii', 'ignore').decode().lower()
        dir_ = unicodedata.normalize('NFD', m.group(3)).encode('ascii', 'ignore').decode().lower()
        for w in PALAVRAS_TERMO:
            if len(esq) >= 2 and len(dir_) >= 2 and w.startswith(esq) and w.endswith(dir_) and len(esq) + len(dir_) in (len(w) - 1, len(w) - 2):
                return True
    return bool(RX_SIGLA_QUEBRADA.search(texto))


class ErroOCR(RuntimeError):
    pass


def recompor_acentos(t, c):
    def cedilha(m):
        c['cedilha'] += 1
        return 'Ç' if m.group(1) == 'C' else 'ç'

    def acento(m):
        c['acentos'] += 1
        return unicodedata.normalize('NFC', m.group(2).replace('ı', 'i') + ACENTOS[m.group(1)])

    t = re.sub(r'([cC])\s?¸', cedilha, t)
    t = re.sub(r'¸\s?([cC])', cedilha, t)
    t = re.sub(r'(?<=[A-Za-zÀ-ÿ]) ?([´`ˆ˜¨]) ?([aeiouAEIOUı])', acento, t)
    t = re.sub(r'([´`ˆ˜¨]) ?([aeiouAEIOUı])', acento, t)
    return t


def limpar(t, c):
    t = re.sub(r'\r\n?', '\n', t)
    for k, v in LIGADURAS.items():
        if k in t:
            c['ligaduras'] += t.count(k)
            t = t.replace(k, v)
    c['hifen_suave'] += t.count('\xad')
    t = t.replace('\xad', '')
    c['nbsp'] += t.count('\xa0')
    t = t.replace('\xa0', ' ')
    t, n = re.subn('[\ue000-\uf8ff]', '•', t)
    c['marcadores_privados'] += n
    t, n = re.subn('[​-‏⁠﻿\x80-\x9f]', '', t)
    c['caracteres_invisiveis_removidos'] += n
    t = recompor_acentos(t, c)
    t = unicodedata.normalize('NFC', t)
    t = re.sub(r'[ \t]{2,}', ' ', t)
    t = re.sub(r' +\n', '\n', t)
    t = re.sub(r'\n{3,}', '\n\n', t)
    return t.strip()


def dehifenizar(t, c, amostra, vocab, compostos):
    def juntar(m):
        antes, a, b = m.group(1), m.group(2), m.group(3)
        junta = (a + b).lower()
        manter = ('/' in antes or 'http' in antes.lower() or b.lower() in CLITICOS or a.lower() in PREFIXOS
                  or (junta not in vocab and (a.lower() in vocab or (a + '-' + b).lower() in compostos)))
        if manter:
            c['hifens_mantidos'] += 1
            if len(amostra) < 30:
                amostra.append('%s-%s (mantido)' % (a, b))
            return antes + a + '-' + b
        c['hifens_unidos'] += 1
        if len(amostra) < 30:
            amostra.append('%s-%s -> %s%s' % (a, b, a, b))
        return antes + a + b
    return re.sub('(\\S*?)([A-Za-zÀ-ÿ]{2,})-\n([a-zà-ÿ]{2,})', juntar, t)


def _ocr_imagem(img):
    base = [TESSERACT, img, 'stdout', '--tessdata-dir', TESSDATA, '-l', 'por', '--psm', OCR_PSM]
    try:
        r_texto = subprocess.run(base, capture_output=True, text=True, encoding='utf-8')
        r_tsv = subprocess.run(base + ['tsv'], capture_output=True, text=True, encoding='utf-8')
    except OSError as e:
        raise ErroOCR('não foi possível executar o Tesseract (%s): %s' % (TESSERACT, e))
    for r in (r_texto, r_tsv):
        if r.returncode != 0:
            raise ErroOCR('Tesseract terminou com código %s: %s' % (r.returncode, (r.stderr or '').strip()[:300]))
    confs = []
    for linha in r_tsv.stdout.splitlines()[1:]:
        col = linha.split('\t')
        if len(col) >= 12 and col[10] not in ('-1', '') and col[11].strip():
            confs.append(float(col[10]))
    return r_texto.stdout, (statistics.mean(confs) if confs else 0.0), len(confs)


def ocr_pagina(pagina):
    melhor = ('', 0.0, 0, 0)
    with tempfile.TemporaryDirectory() as pasta:
        img = os.path.join(pasta, 'p.png')
        for rot in (0, 90, 270):
            pagina.get_pixmap(matrix=pymupdf.Matrix(300 / 72, 300 / 72).prerotate(rot)).save(img)
            txt, conf, n = _ocr_imagem(img)
            if conf * min(n, 50) > melhor[1] * min(melhor[2], 50):
                melhor = (txt, conf, n, rot)
            if rot == 0 and conf >= 80 and n >= 20:
                break
    return melhor[0], melhor[1], melhor[3]


def carregar_mapa_controle():
    mapa = collections.defaultdict(dict)
    if os.path.exists(CONTROLE_CSV):
        with open(CONTROLE_CSV, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                if (r['codigo_curso'] or '').strip():
                    mapa[r['codigo_curso'].strip()][chr(int(r['hex'].strip(), 16))] = r['substituto']
    return mapa


def processar(caminho_pdf, mapa_controle=None):
    nome = os.path.basename(caminho_pdf)
    base = nome[:-4]
    doc = pymupdf.open(caminho_pdf)
    trocas = (mapa_controle or {}).get(nome.split('_')[0], {})
    c = collections.Counter()
    amostra, com_ocr, mistas, saida, paginas, ocrs, pend, por_controle, leves = [], [], [], [], [], [], [], [], []
    for i, p in enumerate(doc, 1):
        area = p.rect.width * p.rect.height
        area_img = 0
        falha_imagem = False
        try:
            for info in p.get_image_info():
                r = pymupdf.Rect(info['bbox']) & p.rect
                area_img += r.width * r.height
        except Exception:
            falha_imagem = True
        bruto = p.get_text()
        for ch, subst in trocas.items():
            if ch in bruto:
                c['caracteres_de_controle_trocados'] += bruto.count(ch)
                bruto = bruto.replace(ch, subst)
        limpo = limpar(bruto, c)
        sobrou = len(RX_CONTROLE.findall(limpo))
        motivo = ''
        if sobrou:
            letras = len(re.findall(r'[A-Za-zÀ-ÿ]', limpo))
            c['paginas_com_controle_sem_mapeamento'] += 1
            c['caracteres_de_controle_removidos'] += sobrou
            if sobrou / (sobrou + letras) > LIMITE_ILEGIVEL:
                motivo = 'texto quase todo ilegível (%d caracteres de controle contra %d letras); conferir no PDF' % (sobrou, letras)
            elif termo_quebrado(limpo):
                motivo = 'termo de busca com caractere de controle dentro (menção possivelmente perdida); conferir no PDF'
            else:
                leves.append(i)
            limpo = RX_CONTROLE.sub('', limpo)
        if falha_imagem and not motivo:
            motivo = 'não foi possível medir as imagens da página (não se sabe se precisa de OCR); conferir no PDF'
        paginas.append((i, limpo, area_img / area if area else 0, motivo))
    todo = '\n'.join(x[1] for x in paginas).lower()
    vocab = set(re.findall(r'(?<![-\w])[a-zà-ÿ]{3,}(?![-\w])', todo))
    compostos = set(re.findall(r'[a-zà-ÿ]{2,}-[a-zà-ÿ]{2,}', todo))
    for i, t, frac_img, motivo_controle in paginas:
        t = dehifenizar(t, c, amostra, vocab, compostos)
        if motivo_controle:
            por_controle.append(i)
            pend.append((i, motivo_controle))
        if len(t) < LIMITE_TEXTO_OCR and frac_img > 0.3:
            com_ocr.append(i)
            try:
                o, conf, rot = ocr_pagina(doc[i - 1])
            except ErroOCR as e:
                raise ErroOCR('%s, página %d: %s' % (nome, i, e))
            o = limpar(o, c)
            o = dehifenizar(o, c, amostra, vocab, compostos)
            if not o.strip():
                pend.append((i, 'OCR rodou mas não leu texto (pode ser página em branco; conferir no PDF)'))
            ocrs.append((i, round(conf), len(o)) if not rot else (i, round(conf), len(o), 'rot%d' % rot))
            saida.append('[[PÁGINA %d]]\n[[OCR confiança=%.0f%%%s]]\n%s' % (i, conf, ' rot%d' % rot if rot else '', o))
        else:
            if frac_img > 0.3:
                mistas.append(i)
            saida.append('[[PÁGINA %d]]\n%s' % (i, t))
    txt = '\n\n'.join(saida) + '\n'
    residuos = {'símbolos_de_acento_restantes': len(re.findall(r'[´`ˆ˜¨¸]', txt)), 'caractere_inválido': txt.count('�') + len(RX_CONTROLE.findall(txt))}
    os.makedirs(TEXTOS_DIR, exist_ok=True)
    arq_pend = os.path.join(PENDENCIAS_DIR, base + '.pendencias.csv')
    if pend:
        os.makedirs(PENDENCIAS_DIR, exist_ok=True)
        with open(arq_pend, 'w', newline='', encoding='utf-8-sig') as f:
            w = csv.writer(f, delimiter=';')
            w.writerow(['arquivo', 'pagina', 'motivo'])
            w.writerows([(nome, pg, motivo) for pg, motivo in pend])
    elif os.path.exists(arq_pend):
        os.remove(arq_pend)
    rel = ['# Relatório de normalização: ' + nome, '', '- Páginas: %d' % len(doc),
           '- Páginas de imagem sem texto (OCR aplicado): %d %s' % (len(com_ocr), com_ocr),
           '- OCR (página, confiança média %, caracteres): ' + str(ocrs) + ' (conferir visualmente as de confiança < 85)',
           '- Páginas com caracteres de controle sem mapeamento: pendência (texto ilegível ou termo quebrado) %d %s; aceitas (poucos símbolos soltos, nenhum termo quebrado) %d %s' % (len(por_controle), por_controle, len(leves), leves),
           '- Páginas com texto e imagem grande (conferir por amostra): %d' % len(mistas), '', '## Correções aplicadas']
    rel += ['- %s: %d' % (k, v) for k, v in sorted(c.items())] or ['- nenhuma']
    rel += ['', '## Resíduos após a normalização (para conferência)'] + ['- %s: %d' % (k, v) for k, v in residuos.items()]
    rel += ['', '## Amostra de hifens tratados'] + ['- ' + a for a in amostra]
    with open(os.path.join(TEXTOS_DIR, base + '_relatorio.md'), 'w', encoding='utf-8', newline='\n') as f:
        f.write('\n'.join(rel) + '\n')
    destino = os.path.join(TEXTOS_DIR, base + '.txt')
    with open(destino + '.tmp', 'w', encoding='utf-8', newline='\n') as f:
        f.write(txt)
    os.replace(destino + '.tmp', destino)
    print('  %s: %d pág | OCR em %d | pendências %d | resíduos de acento %d' % (nome, len(doc), len(com_ocr), len(pend),
                                                                              residuos['símbolos_de_acento_restantes']), flush=True)
    return {'paginas': len(doc), 'ocr': len(com_ocr), 'pendencias': len(pend)}


def mapa_ppc_pdf():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    cab = [c.value for c in next(wb['Dados_Unificados'].iter_rows(max_row=1))]
    i_id, i_cod = cab.index('id_curso'), cab.index('Código do Curso')
    codigo_de = {r[i_id]: str(r[i_cod]).strip() for r in wb['Dados_Unificados'].iter_rows(min_row=2, values_only=True)}
    cab_p = [c.value for c in next(wb['PPC'].iter_rows(max_row=1))]
    j_ppc, j_curso, j_status = cab_p.index('id_ppc'), cab_p.index('id_curso'), cab_p.index('status_coleta')
    linhas = [(r[j_ppc], codigo_de[r[j_curso]]) for r in wb['PPC'].iter_rows(min_row=2, values_only=True) if r[j_status] == 1]
    wb.close()
    por_codigo = {}
    for nome in sorted(os.listdir(PPC_DIR)):
        if nome.lower().endswith('.pdf'):
            por_codigo.setdefault(nome.split('_')[0], []).append(nome)
    saida = []
    for id_ppc, codigo in linhas:
        achados = por_codigo.get(codigo, [])
        if len(achados) != 1:
            raise SystemExit('Curso %s com status_coleta 1 tem %d PDF(s) em entradas/ppc (esperado 1). Rode o Passo 1 de novo.' % (codigo, len(achados)))
        saida.append((id_ppc, codigo, achados[0]))
    return sorted(saida)


def pdfs_para_processar(so=None):
    alvo = [nome for _, _, nome in mapa_ppc_pdf()]
    if so is not None:
        inexistentes = sorted(c for c in so if c not in {a.split('_')[0] for a in alvo})
        if not so or inexistentes:
            raise SystemExit('--so: informe códigos de curso que existam entre os PPCs com PDF (inexistentes: %s)' % (inexistentes or 'nenhum código informado'))
        alvo = [a for a in alvo if a.split('_')[0] in so]
    return alvo


def validar_decisoes(todos, mapa_controle, sem_pdf):
    codigos = {n.split('_')[0] for n in todos}
    erros, de_lado = [], 0
    for c in mapa_controle:
        if c in sem_pdf:
            de_lado += 1
        elif c not in codigos:
            erros.append('caracteres_controle.csv: curso %s não está entre os PPCs com PDF' % c)
    if os.path.exists(CONFERIDAS_CSV):
        with open(CONFERIDAS_CSV, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                c = (r['codigo_curso'] or '').strip()
                if c in sem_pdf:
                    de_lado += 1
                elif c and c not in codigos:
                    erros.append('pendencias_conferidas.csv: curso %s não está entre os PPCs com PDF' % r['codigo_curso'])
    if erros:
        raise SystemExit('Decisões inconsistentes: ' + '; '.join(erros))
    if de_lado:
        print('  decisões de cursos sem PDF em entradas/ppc, deixadas de lado: %d' % de_lado)


def carregar_fora_do_ppc():
    intervalos = collections.defaultdict(list)
    if os.path.exists(FORA_DO_PPC_CSV):
        with open(FORA_DO_PPC_CSV, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                if (r['codigo_curso'] or '').strip():
                    intervalos[r['codigo_curso'].strip()].append((int(r['pag_inicio']), int(r['pag_fim'])))
    return intervalos


def pendencias_abertas(todos):
    conferidas = set()
    if os.path.exists(CONFERIDAS_CSV):
        with open(CONFERIDAS_CSV, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                conferidas.add((r['codigo_curso'].strip(), str(r['pagina']).strip()))
    abertas = []
    for arq in sorted(glob.glob(os.path.join(PENDENCIAS_DIR, '*.pendencias.csv'))):
        with open(arq, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                if r['arquivo'] in todos and (r['arquivo'].split('_')[0], str(r['pagina'])) not in conferidas:
                    abertas.append(r)
    with open(PENDENCIAS_CSV, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=['arquivo', 'pagina', 'motivo'], delimiter=';')
        w.writeheader()
        w.writerows(abertas)
    return abertas


RX_PAGINA = re.compile(r'\[\[PÁGINA (\d+)\]\]\n')
RX_MARCA_OCR = re.compile(r'^\[\[OCR confiança=(\d+)%[^\]]*\]\]\n?')
COLUNAS_TAMANHO = ['n_paginas_pdf', 'n_paginas_validas', 'n_paginas_ocr', 'n_paginas_pendentes', 'n_caracteres']


def paginas_do_texto(caminho):
    with open(caminho, encoding='utf-8') as f:
        partes = RX_PAGINA.split(f.read())
    saida = []
    for i in range(1, len(partes) - 1, 2):
        corpo, conf = partes[i + 1], None
        m = RX_MARCA_OCR.match(corpo)
        if m:
            conf, corpo = int(m.group(1)), corpo[m.end():]
        saida.append((int(partes[i]), conf, len(corpo.strip())))
    return saida


def coletar_estatisticas(mapa, abertas):
    fora = carregar_fora_do_ppc()
    pend_pg = collections.defaultdict(set)
    for r in abertas:
        pend_pg[r['arquivo']].add(int(r['pagina']))
    stats, suspeitas = {}, []
    for id_ppc, codigo, nome in mapa:
        caminho = os.path.join(TEXTOS_DIR, nome[:-4] + '.txt')
        if not os.path.exists(caminho):
            continue
        paginas = paginas_do_texto(caminho)
        dentro = [(pg, conf, n) for pg, conf, n in paginas if not any(a <= pg <= b for a, b in fora.get(codigo, []))]
        pend_dentro = {pg for pg, _, _ in dentro} & pend_pg[nome]
        stats[id_ppc] = [len(paginas), len(dentro) - len(pend_dentro), sum(1 for _, conf, _ in dentro if conf is not None),
                         len(pend_dentro), sum(n for _, _, n in dentro)]
        for pg, conf, n in dentro:
            if conf is not None and (conf < CONFIANCA_SUSPEITA or n < 30):
                suspeitas.append((nome, pg, conf, n))
    return stats, suspeitas


def gravar_tamanho(stats):
    wb = openpyxl.load_workbook(XLSX)
    ws = wb['PPC']
    cab = [c.value for c in ws[1]]
    faltam = [c for c in COLUNAS_TAMANHO if c not in cab]
    if faltam:
        raise SystemExit('A aba PPC não tem as colunas %s: rode o Passo 1 de novo (com --recriar).' % faltam)
    col_ppc = cab.index('id_ppc') + 1
    cols = [cab.index(c) + 1 for c in COLUNAS_TAMANHO]
    for linha in range(2, ws.max_row + 1):
        valores = stats.get(ws.cell(linha, col_ppc).value)
        if valores:
            for col, v in zip(cols, valores):
                ws.cell(linha, col).value = v
    temporario = XLSX + '.tmp'
    wb.save(temporario)
    wb.close()
    os.replace(temporario, XLSX)


def gravar_suspeitas(suspeitas):
    with open(SUSPEITAS_CSV, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['arquivo', 'pagina', 'confianca_ocr', 'caracteres', 'motivo'])
        for nome, pg, conf, n in suspeitas:
            w.writerow([nome, pg, conf, n, 'OCR com confiança baixa' if conf < CONFIANCA_SUSPEITA else 'OCR leu muito pouco texto'])


def executar(reaproveitar=False, so=None):
    if not os.path.exists(XLSX):
        raise SystemExit('Falta %s: rode o Passo 1 antes.' % XLSX)
    os.makedirs(TEXTOS_DIR, exist_ok=True)
    try:
        open(XLSX, 'r+b').close()
    except PermissionError:
        raise SystemExit('O Excel %s está aberto (ou bloqueado). Feche-o antes de rodar este passo: ele grava as colunas de tamanho no fim.' % XLSX)
    mapa = mapa_ppc_pdf()
    todos = [nome for _, _, nome in mapa]
    mapa_controle = carregar_mapa_controle()
    validar_decisoes(todos, mapa_controle, rc.codigos_sem_pdf(XLSX))
    alvo = pdfs_para_processar(so)
    print('  PDFs a processar: %d (%d ao mesmo tempo)' % (len(alvo), PDFS_AO_MESMO_TEMPO), flush=True)
    fazer = []
    for nome in alvo:
        txt_existente = os.path.join(TEXTOS_DIR, nome[:-4] + '.txt')
        if reaproveitar and os.path.exists(txt_existente) and not (
                (os.path.exists(CONTROLE_CSV) and os.path.getmtime(CONTROLE_CSV) > os.path.getmtime(txt_existente))
                or os.path.getmtime(os.path.join(PPC_DIR, nome)) > os.path.getmtime(txt_existente)):
            continue
        fazer.append(nome)
    with concurrent.futures.ProcessPoolExecutor(max_workers=PDFS_AO_MESMO_TEMPO) as ex:
        futuros = [ex.submit(processar, os.path.join(PPC_DIR, nome), mapa_controle) for nome in fazer]
        for f in concurrent.futures.as_completed(futuros):
            try:
                f.result()
            except ErroOCR as e:
                ex.shutdown(cancel_futures=True)
                raise SystemExit('falha no OCR, o passo foi interrompido: %s' % e)
    print('  textos em resultado/textos: %d' % len(glob.glob(os.path.join(TEXTOS_DIR, '*.txt'))))
    abertas = pendencias_abertas(todos)
    stats, suspeitas = coletar_estatisticas(mapa, abertas)
    gravar_suspeitas(suspeitas)
    gravar_tamanho(stats)
    print('  tamanho dos PPCs gravado na aba PPC (%d PPCs) | páginas de OCR suspeito: %d (resultado/paginas_ocr_suspeitas.csv)' % (len(stats), len(suspeitas)))
    if abertas:
        print('  ATENÇÃO: %d página(s) sem texto lido, a conferir (lista em resultado/pendencias_extracao.csv):' % len(abertas))
        for r in abertas[:10]:
            print('    %s p.%s: %s' % (r['arquivo'], r['pagina'], r['motivo']))
        print('  Depois de conferir no PDF, registre em entradas/decisoes/pendencias_conferidas.csv (codigo_curso;pagina;nota).')
        print('  A extração NÃO está completa: os números podem ficar subestimados até as pendências serem conferidas.')
    else:
        print('  Extração completa: nenhuma página sem texto lido pendente.')
    return abertas


if __name__ == '__main__':
    args = sys.argv[1:]
    so = None
    if '--so' in args:
        so = set(a for a in args[args.index('--so') + 1:] if not a.startswith('--'))
    executar(reaproveitar='--reaproveitar' in args, so=so)
