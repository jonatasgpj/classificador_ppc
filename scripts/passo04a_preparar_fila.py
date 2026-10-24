import collections
import csv
import hashlib
import json
import os
import re
import unicodedata

import openpyxl
import pymupdf

pymupdf.TOOLS.mupdf_display_errors(False)

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')
TEXTOS_DIR = os.path.join(AQUI, 'resultado', 'textos')
PPC_DIR = os.path.join(AQUI, 'entradas', 'ppc')
CONGELAMENTO_JSON = os.path.join(AQUI, 'resultado', 'congelamento.json')
SAIDA = os.path.join(AQUI, 'resultado', 'revisao')

JANELA_ASSINATURA = 40
TERMOS = {'Inteligência Artificial': ['inteligencia', 'artificial'], 'Artificial Intelligence': ['artificial', 'intelligence'],
          'IA': ['IA'], 'AI': ['AI']}


def dobra(s):
    return ''.join(unicodedata.normalize('NFD', ch)[0] for ch in s)


def so_alnum(s):
    return re.sub(r'[^a-z0-9]+', '', dobra(s).lower())


def partir_trecho(trecho):
    m = re.search(r'«(.*?)»', trecho, re.S)
    if not m:
        return trecho, '', ''
    return trecho[:m.start()], m.group(1), trecho[m.end():]


def assinatura(termo, trecho):
    antes, _, depois = partir_trecho(trecho)
    texto = so_alnum(antes)[-JANELA_ASSINATURA:] + '|' + so_alnum(termo) + '|' + so_alnum(depois)[:JANELA_ASSINATURA]
    return hashlib.sha1(texto.encode('utf-8')).hexdigest()[:12]


RX_CODIGO = re.compile(r'\b[A-Z]{2,6}[ -]?\d{3,5}[A-Z]?\b')
RX_MATRIZ = re.compile(r'\boptativ|\beletiv|\bperiodo\b|\bsemestre\b|\bsem\.\s|equival|obrigatori|\bch\b|carga hor|creditos|\bt\b p\b', re.I)
RX_MACRO_FORTE = re.compile(r'\b(egressos?|perfil profissional|areas? de atuacao|campo de atuacao|mercado de trabalho|tais como|dentre outras|entre outras)\b')
RX_MACRO_FRACO = re.compile(r'\b(profissional|profissionais|sociedade|mercado|etc)\b')
RX_VERBO = re.compile(r'\b(sao|foi|foram|ser|sera|serao|possui|possuem|permite|visa|busca|oferece|contribui|inclui|abrange|envolve|trata|apresenta|garante|'
                      r'formar|atuar|podem|pode|deve|devem|desenvolver|capacitar|proporcionar|estao|existem|existe|ha|tem)\b')
RX_DOCENTE = re.compile(r'\b(doutor(ado)?|mestre|mestrado|ph\.?d|especialista|lattes|titulacao|dedicacao exclusiva|regime de trabalho|docente)\b')
RX_PESQUISA = re.compile(r'linhas? de pesquisa|grupos? de pesquisa|projetos? de pesquisa')
MICRO = [
    (re.compile(r'\bementa\b|\bementario\b'), 12, 'Ementa'),
    (re.compile(r'\bobjetivos?\b'), 13, 'Objetivos'),
    (re.compile(r'\bcompetencias?\b|\bhabilidades?\b'), 14, 'Competências'),
    (re.compile(r'conteudo programatico|\bconteudos?\b|\bunidades?\b|\bprograma\b'), 15, 'Conteúdos'),
    (re.compile(r'\bmetodologias?\b|procedimentos? didaticos?|estrategias? de ensino'), 16, 'Metodologias'),
    (re.compile(r'\bavaliacao\b|criterios? de avaliacao'), 17, 'Avaliação'),
]
FICHA = re.compile(r'\bementa\b|conteudo programatico|\bdisciplina\s*[:\u2013-]|componente curricular|\bcodigo\s*:|carga[\s-]*horaria|'
                   r'pre-?requisitos?\s*:|\bnatureza\s*:|bibliografia (basica|complementar)|\bobjetivos?\s*:')
INICIO_FICHA = re.compile(r'\bementa\b|\bobjetivos?\b|carga[\s-]*horaria|\bperiodo\b|pre-?requisitos?|\bcreditos?\b|\bcodigo\s*:|\bnatureza\s*:|\bcategoria\s*:|\bch\s*:|\bsemestre\b')
TITULO_FICHA = re.compile(r'\bementa\b|carga[\s-]*horaria|pre-?requisitos?|\bcodigo\s*:|\bnatureza\s*:|\bcategoria\s*:|\bch\s*:|'
                          r'\bobjetivos?\s*:|\bperiodo\s*:|\bcreditos?\s*:')
MACRO_TITULO = [
    (re.compile(r'perfil (do|de) egressos?|perfil profissional'), 4, 'Perfil do egresso'),
    (re.compile(r'\bjustificativa\b'), 2, 'Justificativa'),
    (re.compile(r'objetivos? (geral|gerais|especificos?|do curso)'), 3, 'Objetivos (curso)'),
    (re.compile(r'competencias e habilidades|competencias do egresso'), 5, 'Competências (curso)'),
    (re.compile(r'concepcao (do curso|pedagogica)|principios (norteadores|do curso|pedagogicos)|historico do curso|apresentacao do curso|identificacao do curso'), 1, 'Concepção'),
    (re.compile(r'organizacao curricular|estrutura curricular|matriz curricular|nucleos? (de formacao|basico|especifico)|eixos? (de formacao|tematicos?)|'
                r'atividades complementares|estagio supervisionado'), 6, 'Estrutura curricular'),
]
MACRO_PALAVRA = [
    (re.compile(r'\begressos?\b|campo de atuacao|areas? de atuacao|mercado de trabalho'), 4, 'Perfil do egresso'),
    (re.compile(r'\bdemanda\b|\brelevancia\b|\bnecessidade\b'), 2, 'Justificativa'),
    (re.compile(r'\bobjetivos?\b'), 3, 'Objetivos (curso)'),
    (re.compile(r'\bcompetencias?\b|\bhabilidades?\b'), 5, 'Competências (curso)'),
    (re.compile(r'\bprincipios?\b|\bhistoria\b|\bhistorico\b|\bfilosofia\b'), 1, 'Concepção'),
    (re.compile(r'\bnucleos?\b|\beixos?\b|\btrilhas?\b|\bestagio\b'), 6, 'Estrutura curricular'),
]
MOTIVO_DA_TRIAGEM = {'bibliografia': 'bibliografia', 'sumário': 'sumario', 'fora do PPC': 'fora_do_ppc', 'ata de reunião': 'ata',
                     'tabela de equivalência': 'equivalencia', 'junto com o extenso': 'duplicata_extenso'}


def densidade_num(s):
    pal = s.split()
    return sum(1 for p in pal if re.fullmatch(r'\(?\d+([.,]\d+)?[hH]?\)?', p)) / max(1, len(pal))


def prosa(s):
    pal = s.split()
    return len(RX_VERBO.findall(s)) / max(1, len(pal))


def rotulo_mais_perto(antes_f, tabela, alcance):
    melhor = None
    for rx, idc, nome in tabela:
        for m in rx.finditer(antes_f[-alcance:]):
            dist = len(antes_f[-alcance:]) - m.end()
            if melhor is None or dist < melhor[2]:
                melhor = (idc, nome, dist)
    return melhor


def sugerir(texto_pagina, posicao, comprimento):
    antes = re.sub(r'\s+', ' ', texto_pagina[max(0, posicao - 900):posicao])
    depois = re.sub(r'\s+', ' ', texto_pagina[posicao + comprimento:posicao + comprimento + 500])
    af, df = dobra(antes).lower(), dobra(depois).lower()
    perto = antes[-300:] + ' ⟨IA⟩ ' + depois[:200]
    pf = dobra(perto).lower()
    dn = densidade_num(perto)
    n_cod = len(RX_CODIGO.findall(perto))
    n_mat = len(RX_MATRIZ.findall(dobra(perto)))
    if len(RX_DOCENTE.findall(pf)) >= 2:
        return None, 'sem', None, 'sem', '', 'parece lista/quadro de docentes'
    if RX_PESQUISA.search(pf):
        return None, 'sem', None, 'sem', '', 'parece linha/grupo de pesquisa'
    if re.search(r'laboratorio', af[-60:]):
        return None, 'sem', None, 'sem', '', 'parece laboratório/infraestrutura'
    if TITULO_FICHA.search(df[:80]) and dn < 0.25:
        rot = rotulo_mais_perto(af, MICRO, 100)
        if rot:
            return 'Micro', 'fraca', rot[0], 'fraca', rot[1], ''
        return 'Micro', 'forte', 12, 'forte', 'título da ficha', ''
    rot = rotulo_mais_perto(af, MICRO, 500)
    if rot and FICHA.search(af[-500:]) and dn < 0.12:
        forca = 'forte' if rot[2] <= 150 else 'fraca'
        return 'Micro', forca, rot[0], forca, rot[1], ''
    if FICHA.search(af[-400:]) and dn < 0.12:
        return 'Micro', 'fraca', 12, 'fraca', 'ficha (sem rótulo de campo)', ''
    if dn >= 0.12 or n_mat >= 2 or n_cod >= 2:
        forte = dn >= 0.2 and n_cod >= 1
        sem_codigo_perto = not RX_CODIGO.search(antes[-80:])
        if re.search(r'pre-?requisitos?\s*:?', af[-60:]):
            return 'Meso', 'forte', 11, 'forte', 'pré-requisito', ''
        if re.search(r'carga[\s-]*horaria|total de horas|resumo', af[-100:]) and sem_codigo_perto:
            return 'Meso', 'fraca', 8, 'fraca', 'carga horária', ''
        if re.search(r'fluxograma|sequencia|\btrilha', af[-120:]) and sem_codigo_perto:
            return 'Meso', 'fraca', 9, 'fraca', 'sequência', ''
        if re.search(r'projeto integrador|integracao|interdisciplin', af[-60:]) and sem_codigo_perto:
            return 'Meso', 'fraca', 10, 'fraca', 'integração', ''
        f = 'forte' if forte else 'fraca'
        return 'Meso', f, 7, f, 'componente na matriz', ''
    forte_macro = RX_MACRO_FORTE.search(pf)
    if (forte_macro or RX_MACRO_FRACO.search(pf) or prosa(pf) >= 0.025) and dn < 0.06 and (forte_macro or prosa(pf) >= 0.025):
        forca = 'forte' if forte_macro else 'fraca'
        rot = rotulo_mais_perto(af, MACRO_TITULO, 400)
        if rot:
            return 'Macro', forca, rot[0], ('forte' if rot[2] <= 300 else 'fraca'), rot[1], ''
        rot = rotulo_mais_perto(af, MACRO_PALAVRA, 600)
        if rot:
            return 'Macro', forca, rot[0], 'fraca', rot[1], ''
        return 'Macro', forca, None, 'sem', '', ''
    return None, 'sem', None, 'sem', '', ''


def compacto(s, sigla=False):
    if sigla:
        return re.sub(r'[^A-Za-z0-9]+', '', s)
    s = ''.join(ch for ch in unicodedata.normalize('NFKD', s) if not unicodedata.combining(ch))
    return re.sub(r'[^a-z0-9]+', '', s.lower())


def caixas_do_termo(pagina_pdf, termo, n_esperado=None, trocas=None):
    sigla = termo in ('IA', 'AI')
    palavras = pagina_pdf.get_text('words')
    for k, (ch, subst) in enumerate(list((trocas or {}).items())):
        palavras = [(*p[:4], p[4].replace(ch, subst), *p[5:]) for p in palavras]
    if sigla:
        return [[tuple(p[:4])] for p in palavras if re.sub(r'^[\W_]+|[\W_]+$', '', p[4]) == termo]
    alvo = compacto(termo)
    alvos = [alvo] + ([alvo.replace(subst, '') for subst in (trocas or {}).values() if subst.isalpha() and subst in alvo] or [])
    toks = [compacto(p[4]) for p in palavras]
    por_janela, i = [], 0
    while i < len(palavras):
        achou = False
        for tam in range(1, 5):
            janela = ''.join(toks[i:i + tam])
            if any(a in janela and (tam == 1 or a not in ''.join(toks[i:i + tam - 1])) and a not in ''.join(toks[i + 1:i + tam]) for a in alvos):
                por_janela.append([tuple(palavras[i + k][:4]) for k in range(tam)])
                i += tam
                achou = True
                break
        if not achou:
            i += 1
    if n_esperado is None or len(por_janela) == n_esperado:
        return por_janela
    try:
        por_busca = [[tuple(r)] for r in pagina_pdf.search_for(termo, hit_max=200)]
    except TypeError:
        por_busca = [[tuple(r)] for r in pagina_pdf.search_for(termo)]
    if len(por_busca) == n_esperado:
        return por_busca
    return por_janela if len(por_janela) >= len(por_busca) else por_busca


def carregar_mapa_controle():
    caminho = os.path.join(AQUI, 'entradas', 'decisoes', 'caracteres_controle.csv')
    mapa = collections.defaultdict(dict)
    if os.path.exists(caminho):
        with open(caminho, encoding='utf-8-sig', newline='') as f:
            for r in csv.DictReader(f, delimiter=';'):
                if (r['codigo_curso'] or '').strip():
                    mapa[r['codigo_curso'].strip()][chr(int(r['hex'].strip(), 16))] = r['substituto']
    return mapa


def calcular_destaques(itens):
    docs = {}
    mapa_controle = carregar_mapa_controle()
    por_pagina = collections.defaultdict(list)
    saida = {}
    for it in itens:
        if it['posicao'] < 0:
            saida[it['chave']] = {'estado': 'manual', 'caixas': []}
        else:
            por_pagina[(it['arquivo'], it['pagina'], it['termo'])].append(it)
    for (arquivo, pagina, termo), linhas in por_pagina.items():
        linhas.sort(key=lambda x: x['posicao'])
        if arquivo not in docs:
            docs[arquivo] = pymupdf.open(os.path.join(PPC_DIR, arquivo))
        pg = docs[arquivo][pagina - 1]
        w, h = pg.rect.width, pg.rect.height
        ocr = linhas[0]['pagina_ocr']
        achados = [] if ocr else caixas_do_termo(pg, termo, len(linhas), mapa_controle.get(arquivo.split('_')[0]))
        todas = []
        for oc in achados:
            caixas = []
            for c in oc:
                r = pymupdf.Rect(c) * pg.rotation_matrix
                r.normalize()
                caixas.append([round(r.x0 / w, 4), round(r.y0 / h, 4), round(r.x1 / w, 4), round(r.y1 / h, 4)])
            todas.append(caixas)
        exato = (not ocr) and len(achados) == len(linhas)
        for rank, it in enumerate(linhas):
            if ocr:
                saida[it['chave']] = {'estado': 'escaneada', 'caixas': []}
            elif exato:
                saida[it['chave']] = {'estado': 'ok', 'caixas': [{'caixas': oc, 'atual': int(i == rank)} for i, oc in enumerate(todas)]}
            else:
                saida[it['chave']] = {'estado': 'ambiguo' if achados else 'sem_palavras',
                                      'caixas': [{'caixas': oc, 'atual': 0} for oc in todas]}
    for d in docs.values():
        d.close()
    return saida


def ler_dados():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    cab_c = [c.value for c in next(wb['Dados_Unificados'].iter_rows(max_row=1))]
    codigo_de_curso = {r[cab_c.index('id_curso')]: str(r[cab_c.index('Código do Curso')]).strip() for r in wb['Dados_Unificados'].iter_rows(min_row=2, values_only=True)}
    sigla_de_curso = {r[cab_c.index('id_curso')]: r[cab_c.index('Sigla da IES')] for r in wb['Dados_Unificados'].iter_rows(min_row=2, values_only=True)}
    ppcs = {}
    for r in wb['PPC'].iter_rows(min_row=2, values_only=True):
        ppcs[r[0]] = {'id_ppc': r[0], 'id_curso': r[1], 'codigo_curso': codigo_de_curso[r[1]], 'sigla': sigla_de_curso[r[1]], 'status': r[5]}
    cab_o = [c.value for c in next(wb['OCORRENCIA'].iter_rows(max_row=1))]
    ocorrencias = [dict(zip(cab_o, r)) for r in wb['OCORRENCIA'].iter_rows(min_row=2, values_only=True)]
    wb.close()
    arquivos = {}
    for nome in os.listdir(PPC_DIR):
        if nome.lower().endswith('.pdf'):
            arquivos[nome.split('_')[0]] = nome
    return ppcs, ocorrencias, arquivos


def paginas_com_ocr(caminho):
    with open(caminho, encoding='utf-8') as f:
        partes = re.split(r'\[\[PÁGINA (\d+)\]\]\n', f.read())
    saida = {}
    for i in range(1, len(partes) - 1, 2):
        corpo = partes[i + 1]
        m = re.match(r'^\[\[OCR confiança=(\d+)%[^\]]*\]\]\n?', corpo)
        if m:
            corpo = corpo[m.end():]
        saida[int(partes[i])] = (corpo, bool(m))
    return saida


def verificar_congelamento(ocorrencias):
    if not os.path.exists(CONGELAMENTO_JSON):
        raise SystemExit('Falta resultado/congelamento.json: rode o Passo 3.')
    ref = json.load(open(CONGELAMENTO_JSON, encoding='utf-8'))
    atual = hashlib.sha256(json.dumps([[o[k] for k in ('id_ocorrencia', 'id_ppc', 'pagina', 'posicao', 'termo', 'trecho', 'triagem_auto', 'ocr_fraco')]
                                      for o in ocorrencias], ensure_ascii=False).encode('utf-8')).hexdigest()
    if atual != ref['ocorrencia'] or len(ocorrencias) != ref['n_ocorrencias']:
        raise SystemExit('A aba OCORRENCIA mudou depois do Passo 3 (hash diferente de congelamento.json). Rode o Passo 3 de novo.')
    for nome, caminho in (('paginas_fora_do_ppc', os.path.join(AQUI, 'entradas', 'decisoes', 'paginas_fora_do_ppc.csv')),
                          ('termos', os.path.join(AQUI, 'entradas', 'termos.csv')),
                          ('mencoes_manuais', os.path.join(AQUI, 'entradas', 'decisoes', 'mencoes_manuais.csv'))):
        atual_arq = None
        if os.path.exists(caminho):
            with open(caminho, 'rb') as f:
                atual_arq = hashlib.sha256(f.read()).hexdigest()
        if atual_arq != ref.get(nome):
            raise SystemExit('%s mudou depois do Passo 3 (hash diferente de congelamento.json). Rode o Passo 3 de novo.' % nome)
    for base, h in ref['textos'].items():
        caminho = os.path.join(TEXTOS_DIR, base + '.txt')
        with open(caminho, 'rb') as f:
            if hashlib.sha256(f.read()).hexdigest() != h:
                raise SystemExit('O texto %s mudou depois do Passo 3. Rode o Passo 3 de novo (e refaça a fila).' % base)
    return ref


def montar():
    ppcs, ocorrencias, arquivos = ler_dados()
    verificar_congelamento(ocorrencias)
    itens = []
    paginas_cache = {}
    for o in ocorrencias:
        p = ppcs[o['id_ppc']]
        arquivo = arquivos[p['codigo_curso']]
        base = arquivo[:-4]
        if base not in paginas_cache:
            paginas_cache[base] = paginas_com_ocr(os.path.join(TEXTOS_DIR, base + '.txt'))
        texto, ocr = paginas_cache[base][o['pagina']]
        _, meio, _ = partir_trecho(o['trecho'])
        if o['posicao'] < 0:
            nivel, forca, id_topico, forca_topico, rotulo, aviso = None, 'sem', None, 'sem', '', ''
        else:
            nivel, forca, id_topico, forca_topico, rotulo, aviso = sugerir(texto, o['posicao'], len(meio))
        itens.append({
            'chave': '%s|%d|%d' % (p['codigo_curso'], o['pagina'], o['posicao']), 'assinatura': assinatura(o['termo'], o['trecho']),
            'id_ocorrencia': o['id_ocorrencia'], 'id_ppc': o['id_ppc'], 'codigo_curso': p['codigo_curso'], 'sigla': p['sigla'], 'arquivo': arquivo,
            'pagina': o['pagina'], 'posicao': o['posicao'], 'termo': o['termo'], 'trecho': o['trecho'], 'triagem_auto': o['triagem_auto'],
            'ocr_fraco': o['ocr_fraco'], 'pagina_ocr': ocr,
            'sug_nivel': nivel, 'sug_forca': forca, 'sug_topico': id_topico, 'sug_forca_topico': forca_topico, 'sug_rotulo': rotulo, 'sug_aviso': aviso,
            'sug_motivo': MOTIVO_DA_TRIAGEM.get(o['triagem_auto'], ''),
        })
    chaves = [i['chave'] for i in itens]
    if len(chaves) != len(set(chaves)):
        raise SystemExit('Chave codigo_curso|pagina|posicao repetida: erro na fila.')
    return itens


def executar():
    itens = montar()
    destaques = calcular_destaques(itens)
    os.makedirs(SAIDA, exist_ok=True)
    escrever_json('fila.json', itens)
    escrever_json('destaques.json', destaques)
    with open(os.path.join(SAIDA, 'sugestoes.csv'), 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.writer(f, delimiter=';')
        w.writerow(['chave', 'assinatura', 'sug_nivel', 'sug_forca', 'sug_topico', 'sug_forca_topico', 'sug_rotulo', 'sug_aviso', 'sug_motivo'])
        for i in itens:
            w.writerow([i['chave'], i['assinatura'], i['sug_nivel'] or '', i['sug_forca'], i['sug_topico'] or '', i['sug_forca_topico'], i['sug_rotulo'],
                        i['sug_aviso'], i['sug_motivo']])
    ref = json.load(open(CONGELAMENTO_JSON, encoding='utf-8'))
    sha = {}
    for nome in ('fila.json', 'destaques.json', 'sugestoes.csv'):
        with open(os.path.join(SAIDA, nome), 'rb') as f:
            sha[nome] = hashlib.sha256(f.read()).hexdigest()
    escrever_json('manifesto.json', {'sha256_ocorrencia': ref['ocorrencia'], 'sha256': sha, 'n_itens': len(itens)})
    resumo(itens, destaques)


def escrever_json(nome, dados):
    with open(os.path.join(SAIDA, nome), 'w', encoding='utf-8', newline='\n') as f:
        json.dump(dados, f, ensure_ascii=False, indent=0)


def resumo(itens, destaques):
    mantidas = [i for i in itens if i['triagem_auto'] == 'mantida']
    print('  ocorrências: %d (mantidas pela triagem %d, descartadas %d)' % (len(itens), len(mantidas), len(itens) - len(mantidas)))
    print('  classificação automática de nível (mantidas): %s' % dict(collections.Counter((i['sug_nivel'] or 'sem') + '/' + i['sug_forca'] for i in mantidas)))
    print('  destaque no PDF: %s' % dict(collections.Counter(d['estado'] for d in destaques.values())))
    print('  gravado em ' + SAIDA)


if __name__ == '__main__':
    executar()
