import hashlib
import io
import json
import os
import sys
import threading
import urllib.parse
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer

import openpyxl
import pymupdf
from PIL import Image, ImageDraw

pymupdf.TOOLS.mupdf_display_errors(False)

AQUI = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(AQUI, 'scripts'))
import passo04a_preparar_fila as fila_mod
import revisao_comum as rc

XLSX = os.path.join(AQUI, 'resultado', 'Dados_Unificados.xlsx')
PPC_DIR = os.path.join(AQUI, 'entradas', 'ppc')
REV = os.path.join(AQUI, 'resultado', 'revisao')
CACHE = os.path.join(REV, 'cache_paginas')
TRAVA = os.path.join(REV, 'servidor.lock')
TELA = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'revisao_tela.html')
LEITOR = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'revisao_leitor.html')
PDFJS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'pdfjs')
LOTE_TODOS = 'todos'
MAX_CORPO = 20000
trava = threading.Lock()


def carregar():
    for nome in ('fila.json', 'destaques.json', 'manifesto.json'):
        if not os.path.exists(os.path.join(REV, nome)):
            raise SystemExit('Falta resultado/revisao/%s: rode o rodar_tudo.py (ou o passo 4A).' % nome)
    itens = json.load(open(os.path.join(REV, 'fila.json'), encoding='utf-8'))
    destaques = json.load(open(os.path.join(REV, 'destaques.json'), encoding='utf-8'))
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    categorias = [dict(zip(['id', 'nivel', 'unidade', 'nome', 'tipo'], r)) for r in wb['CATEGORIA'].iter_rows(min_row=2, values_only=True)]
    wb.close()
    return itens, destaques, categorias


def ocorrencias_do_excel():
    wb = openpyxl.load_workbook(XLSX, read_only=True)
    cab = [c.value for c in next(wb['OCORRENCIA'].iter_rows(max_row=1))]
    saida = [dict(zip(cab, r)) for r in wb['OCORRENCIA'].iter_rows(min_row=2, values_only=True)]
    wb.close()
    return saida


class Revisao:
    def __init__(self):
        self.itens, self.destaques, self.categorias = carregar()
        self.por_chave = {i['chave']: i for i in self.itens}
        self.cats = {c['id']: {'nivel': c['nivel'], 'tipo': c['tipo']} for c in self.categorias}
        self.lotes = self._montar_lotes()
        original = rc.ler_tabela()
        self.classificacao, religadas, orfas = rc.religar(original, self.itens)
        self.orfas = {c: original[c] for c in orfas}
        self.historico = []
        if religadas:
            self._gravar()
            print('  %d linha(s) da tabela religada(s) pela assinatura e gravada(s) com a chave atual.' % len(religadas))
        if orfas:
            print('  ATENÇÃO: %d linha(s) da tabela sem ocorrência correspondente na lista (ex.: %s).' % (len(orfas), orfas[:3]))

    def _montar_lotes(self):
        ordem = sorted(self.itens, key=lambda i: i['id_ocorrencia'])
        lotes = {LOTE_TODOS: {'lote': LOTE_TODOS, 'nome': 'TODOS', 'itens': [i['chave'] for i in ordem]}}
        for i in ordem:
            lid = 'pdf|' + i['arquivo']
            lotes.setdefault(lid, {'lote': lid, 'nome': i['arquivo'][:-4], 'itens': []})['itens'].append(i['chave'])
        return lotes

    def _gravar(self):
        linhas = dict(self.orfas)
        linhas.update(self.classificacao)
        rc.gravar_tabela(linhas)

    def respondido(self, chave):
        reg = self.classificacao.get(chave)
        return bool(reg) and reg['decisao'] in rc.FINAIS

    def resumo(self):
        return [{'lote': l['lote'], 'nome': l['nome'], 'total': len(l['itens']), 'respondidos': sum(1 for c in l['itens'] if self.respondido(c))}
                for l in self.lotes.values()]

    def item_para_tela(self, chave):
        i = self.por_chave[chave]
        saida = {k: i[k] for k in ('chave', 'sigla', 'arquivo', 'pagina', 'termo', 'trecho', 'ocr_fraco', 'triagem_auto',
                                   'sug_nivel', 'sug_forca', 'sug_topico', 'sug_rotulo', 'sug_aviso', 'sug_motivo')}
        saida['destaque'] = self.destaques[chave]['estado']
        saida['resposta'] = self.classificacao.get(chave)
        return saida

    def registrar(self, ev):
        chave = ev.get('chave')
        if chave not in self.por_chave:
            return 'chave desconhecida'
        if ev.get('acao') == 'desfazer':
            k = next((n for n in range(len(self.historico) - 1, -1, -1) if self.historico[n][0] == chave), None)
            if k is None:
                return 'nada para desfazer nesta ocorrência (desde que o servidor foi aberto)'
            _, anterior = self.historico.pop(k)
            if anterior is None:
                self.classificacao.pop(chave, None)
            else:
                self.classificacao[chave] = anterior
            self._gravar()
            return None
        if ev.get('acao') != 'responder':
            return 'ação inválida'
        lote = self.lotes.get(ev.get('lote'))
        if lote is None:
            return 'lista desconhecida'
        if chave not in lote['itens']:
            return 'esta ocorrência não pertence a essa lista'
        d = ev.get('decisao')
        item = self.por_chave[chave]
        anterior = self.classificacao.get(chave) or {}
        texto = lambda campo: (anterior.get(campo) or '') if ev.get(campo) is None else str(ev.get(campo)).strip()
        reg = {'chave': chave, 'assinatura': item['assinatura'], 'termo': item['termo'], 'decisao': d,
               'nivel': ev.get('nivel') if d in ('conta', 'sem_topico') else '',
               'id_categoria': str(ev.get('id_categoria')) if d == 'conta' else '',
               'motivo': (ev.get('motivo') or '') if d in ('conta', 'descarta') else '',
               'componente': texto('componente'), 'nota': texto('nota')}
        erro = rc.validar(reg, self.cats)
        if erro:
            return erro
        self.historico.append((chave, self.classificacao.get(chave)))
        self.classificacao[chave] = reg
        self._gravar()
        return None


def imagem_pagina(rev, chave, deslocamento, dpi):
    item = rev.por_chave[chave]
    arquivo, pagina = item['arquivo'], item['pagina'] + deslocamento
    if deslocamento == 0:
        h = hashlib.sha1(json.dumps(rev.destaques[chave], sort_keys=True).encode('utf-8')).hexdigest()[:10]
        sufixo = 'h' + h
    else:
        sufixo = 'n'
    alvo = os.path.join(CACHE, '%s_p%d_%d_%s.jpg' % (arquivo[:-4], pagina, dpi, sufixo))
    try:
        with open(alvo, 'rb') as f:
            return f.read()
    except OSError:
        pass
    doc = pymupdf.open(os.path.join(PPC_DIR, arquivo))
    try:
        if pagina < 1 or pagina > len(doc):
            return None
        pix = doc[pagina - 1].get_pixmap(matrix=pymupdf.Matrix(dpi / 72, dpi / 72))
    finally:
        doc.close()
    img = Image.frombytes('RGB', (pix.width, pix.height), pix.samples).convert('RGBA')
    if deslocamento == 0:
        sobre = Image.new('RGBA', img.size, (0, 0, 0, 0))
        pincel = ImageDraw.Draw(sobre)
        for oc in rev.destaques[chave]['caixas']:
            cor = (255, 120, 0, 110) if oc['atual'] else (255, 230, 0, 90)
            for x0, y0, x1, y1 in oc['caixas']:
                pincel.rectangle([x0 * pix.width - 1, y0 * pix.height - 1, x1 * pix.width + 1, y1 * pix.height + 1], fill=cor)
        img = Image.alpha_composite(img, sobre)
    os.makedirs(CACHE, exist_ok=True)
    buf = io.BytesIO()
    img.convert('RGB').save(buf, 'JPEG', quality=78)
    tmp = '%s.%d.tmp' % (alvo, threading.get_ident())
    with open(tmp, 'wb') as f:
        f.write(buf.getvalue())
    try:
        os.replace(tmp, alvo)
    except OSError:
        try:
            os.remove(tmp)
        except OSError:
            pass
    return buf.getvalue()


class Manipulador(BaseHTTPRequestHandler):
    rev = None
    porta = 8765

    def _hosts_ok(self):
        host = (self.headers.get('Host') or '').lower()
        return host in ('127.0.0.1:%d' % self.porta, 'localhost:%d' % self.porta)

    def _origem_ok(self):
        origem = self.headers.get('Origin')
        return origem is None or origem.lower() in ('http://127.0.0.1:%d' % self.porta, 'http://localhost:%d' % self.porta)

    def _enviar(self, codigo, corpo, tipo):
        if isinstance(corpo, str):
            corpo = corpo.encode('utf-8')
        self.send_response(codigo)
        self.send_header('Content-Type', tipo)
        self.send_header('Content-Length', str(len(corpo)))
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(corpo)

    def _json(self, dados, codigo=200):
        self._enviar(codigo, json.dumps(dados, ensure_ascii=False), 'application/json; charset=utf-8')

    def log_message(self, *a):
        pass

    def do_GET(self):
        if not self._hosts_ok():
            return self._enviar(403, 'host não permitido', 'text/plain; charset=utf-8')
        url = urllib.parse.urlparse(self.path)
        rota, q = url.path, urllib.parse.parse_qs(url.query)
        rev = self.rev
        try:
            if rota == '/':
                with open(TELA, encoding='utf-8') as f:
                    return self._enviar(200, f.read(), 'text/html; charset=utf-8')
            if rota == '/api/estado':
                return self._json({'lotes': rev.resumo(), 'motivos': rc.MOTIVOS, 'tipos_ementa': rc.TIPOS_EMENTA, 'categorias': rev.categorias,
                                   'total_respondidas': sum(1 for c in rev.por_chave if rev.respondido(c)), 'total_itens': len(rev.itens)})
            if rota.startswith('/api/lote/'):
                lote = rev.lotes.get(urllib.parse.unquote(rota[len('/api/lote/'):]))
                if not lote:
                    return self._json({'erro': 'lista desconhecida'}, 404)
                return self._json({'lote': lote['lote'], 'nome': lote['nome'], 'itens': [{'chave': c, 'respondido': rev.respondido(c)} for c in lote['itens']]})
            if rota == '/leitor':
                with open(LEITOR, encoding='utf-8') as f:
                    return self._enviar(200, f.read(), 'text/html; charset=utf-8')
            if rota.startswith('/pdfjs/'):
                nome = rota[len('/pdfjs/'):]
                if nome not in ('pdf.min.mjs', 'pdf.worker.min.mjs'):
                    return self._enviar(404, 'arquivo não permitido', 'text/plain; charset=utf-8')
                with open(os.path.join(PDFJS, nome), 'rb') as f:
                    return self._enviar(200, f.read(), 'text/javascript; charset=utf-8')
            if rota == '/api/destaque':
                chave = q.get('chave', [''])[0]
                if chave not in rev.destaques:
                    return self._json({'erro': 'chave desconhecida'}, 404)
                return self._json(rev.destaques[chave])
            if rota == '/api/item':
                chave = q.get('chave', [''])[0]
                if chave not in rev.por_chave:
                    return self._json({'erro': 'chave desconhecida'}, 404)
                return self._json(rev.item_para_tela(chave))
            if rota == '/pagina':
                chave = q.get('chave', [''])[0]
                if chave not in rev.por_chave:
                    return self._enviar(404, 'chave desconhecida', 'text/plain; charset=utf-8')
                desloc = max(-1, min(1, int(q.get('d', ['0'])[0])))
                dados = imagem_pagina(rev, chave, desloc, 110)
                if dados is None:
                    return self._enviar(404, 'página fora do PDF', 'text/plain; charset=utf-8')
                return self._enviar(200, dados, 'image/jpeg')
            if rota.startswith('/pdf/'):
                nome = urllib.parse.unquote(rota[len('/pdf/'):])
                if nome not in {i['arquivo'] for i in rev.itens}:
                    return self._enviar(404, 'arquivo não permitido', 'text/plain; charset=utf-8')
                with open(os.path.join(PPC_DIR, nome), 'rb') as f:
                    return self._enviar(200, f.read(), 'application/pdf')
            return self._enviar(404, 'não encontrado', 'text/plain; charset=utf-8')
        except Exception as e:
            return self._json({'erro': 'falha no servidor: %r' % e}, 500)

    def do_POST(self):
        if self.path != '/api/evento':
            return self._enviar(404, 'não encontrado', 'text/plain; charset=utf-8')
        if not self._hosts_ok() or not self._origem_ok():
            return self._json({'ok': False, 'erro': 'origem não permitida'}, 403)
        if (self.headers.get('Content-Type') or '').split(';')[0].strip().lower() != 'application/json':
            return self._json({'ok': False, 'erro': 'Content-Type precisa ser application/json'}, 415)
        try:
            tamanho = int(self.headers.get('Content-Length', '0'))
            if tamanho <= 0 or tamanho > MAX_CORPO:
                return self._json({'ok': False, 'erro': 'corpo inválido'}, 400)
            ev = json.loads(self.rfile.read(tamanho).decode('utf-8'))
            if not isinstance(ev, dict):
                return self._json({'ok': False, 'erro': 'evento inválido'}, 400)
            with trava:
                erro = self.rev.registrar(ev)
            if erro:
                return self._json({'ok': False, 'erro': erro}, 400)
            return self._json({'ok': True, 'total_respondidas': sum(1 for c in self.rev.por_chave if self.rev.respondido(c))})
        except Exception as e:
            return self._json({'ok': False, 'erro': 'falha no servidor: %r' % e}, 500)


class Servidor(ThreadingHTTPServer):
    allow_reuse_address = False


def processo_vivo(pid):
    if pid <= 0:
        return False
    if os.name == 'nt':
        import ctypes
        kernel = ctypes.windll.kernel32
        h = kernel.OpenProcess(0x1000, False, pid)
        if not h:
            return False
        codigo = ctypes.c_ulong()
        ok = kernel.GetExitCodeProcess(h, ctypes.byref(codigo))
        kernel.CloseHandle(h)
        return bool(ok) and codigo.value == 259
    try:
        os.kill(pid, 0)
        return True
    except OSError:
        return False


def adquirir_trava():
    os.makedirs(REV, exist_ok=True)
    if os.path.exists(TRAVA):
        try:
            with open(TRAVA, encoding='utf-8') as f:
                antigo = int((f.read() or '0').strip() or 0)
        except (OSError, ValueError):
            antigo = 0
        if not processo_vivo(antigo):
            os.remove(TRAVA)
    try:
        fd = os.open(TRAVA, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
    except FileExistsError:
        raise SystemExit('Já há um servidor da revisão rodando (ou ele caiu e deixou %s). Feche o outro, ou apague esse arquivo se tiver certeza de que não há outro.' % TRAVA)
    os.write(fd, str(os.getpid()).encode())
    os.close(fd)


def executar(porta=8765, abrir=False):
    if not os.path.exists(XLSX):
        raise SystemExit('Falta %s: rode o rodar_tudo.py.' % XLSX)
    fila_mod.verificar_congelamento(ocorrencias_do_excel())
    try:
        rc.verificar_fila_atual(REV, fila_mod.CONGELAMENTO_JSON)
    except rc.ErroFila as e:
        raise SystemExit('Lista de ocorrências desatualizada: %s' % e)
    adquirir_trava()
    try:
        try:
            Manipulador.rev = Revisao()
        except rc.ErroTabela as e:
            raise SystemExit('Tabela de classificação com problema: %s' % e)
        Manipulador.porta = porta
        servidor = Servidor(('127.0.0.1', porta), Manipulador)
        r = Manipulador.rev
        print('  Pronto: abra http://127.0.0.1:%d no navegador (Ctrl+C para encerrar).' % porta)
        print('  Ocorrências: %d | classificadas: %d | listas: TODOS e %d PDFs | tabela: %s' % (
            len(r.itens), sum(1 for c in r.por_chave if r.respondido(c)), len(r.lotes) - 1, rc.TABELA))
        if abrir:
            webbrowser.open('http://127.0.0.1:%d' % porta)
        try:
            servidor.serve_forever()
        except KeyboardInterrupt:
            print('\n  Encerrado. A tabela já está gravada.')
    finally:
        try:
            os.remove(TRAVA)
        except OSError:
            pass


if __name__ == '__main__':
    executar(int(sys.argv[sys.argv.index('--porta') + 1]) if '--porta' in sys.argv else 8765, abrir='--abrir' in sys.argv)
